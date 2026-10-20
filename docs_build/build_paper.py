"""Render paper_content to (1) LaTeX with the official IEEEtran class -> PDF via Tectonic, (2) a Word copy, (3) a source zip."""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
import paper_content as P  # noqa: E402
from common import FIG, OUT, PREFIX, ROOT  # noqa: E402
from docx_helpers import roman, set_cell_borders, to_pdf_word  # noqa: E402

BUILD = ROOT / "docs_build" / "paper_tex"
CLS = ROOT / "tools" / "dl" / "ieeetran" / "IEEEtran" / "IEEEtran.cls"
TECTONIC = ROOT / "tools" / "tectonic" / "tectonic.exe"

SECS = P.sections()
REFS = {p["id"]: p for p in P.references()}
ORDER: list[str] = []
for s in SECS:
    for para in s["p"]:
        for grp in re.findall(r"\[\[cite:([^\]]+)\]\]", para):
            for k in grp.split(","):
                if k not in ORDER:
                    ORDER.append(k)
NUM = {k: i + 1 for i, k in enumerate(ORDER)}


# ------------------------------------------------------------------ LaTeX
def tex_escape(t: str) -> str:
    t = t.replace("\\", r"\textbackslash{}")
    for a, b in [("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}"),
                 ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}")]:
        t = t.replace(a, b)
    return t


def tex_para(t: str) -> str:
    cites = {}

    def keep(m):
        key = f"CITE{len(cites)}X"
        cites[key] = r"~\cite{" + m.group(1) + "}"
        return key
    t = re.sub(r"\s*\[\[cite:([^\]]+)\]\]", keep, t)
    t = tex_escape(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", t)
    t = re.sub(r"\*(.+?)\*", r"\\emph{\1}", t)
    t = t.replace("\u2264", r"$\le$").replace(" r = ", " $r$ = ")
    for k, v in cites.items():
        t = t.replace(k, v)
    return t


def bib_entry(p: dict) -> str:
    venue = tex_escape(p.get("venue") or "")
    s = f"{tex_escape(p.get('authors', ''))}, ``{tex_escape(p.get('title', ''))},'' {venue}"
    if p.get("doi") and p["doi"] not in (p.get("venue") or ""):
        s += f", doi: {tex_escape(p['doi'])}"
    s = s.rstrip(". ") + "."
    if p.get("url") and not p.get("doi"):
        s += r" [Online]. Available: \url{" + p["url"] + "}"
    return s


def build_tex() -> str:
    L = [r"\documentclass[conference,a4paper]{IEEEtran}", r"\IEEEoverridecommandlockouts",
         r"\usepackage{cite}", r"\usepackage{graphicx}", r"\usepackage{booktabs}", r"\usepackage{textcomp}",
         r"\usepackage[hyphens]{url}", r"\usepackage[T1]{fontenc}", r"\usepackage{newtxtext,newtxmath}",
         r"\def\BibTeX{{\rm B\kern-.05em{\sc i\kern-.025em b}\kern-.08em T\kern-.1667em\lower.7ex\hbox{E}\kern-.125emX}}",
         r"\begin{document}", r"\title{" + tex_escape(P.TITLE) + "}", r"\author{"]
    blocks = []
    for name, reg, mail in P.AUTHORS:
        blocks.append(r"\IEEEauthorblockN{" + tex_escape(name) + " (" + reg + ")}" + "\n" +
                      r"\IEEEauthorblockA{\textit{" + tex_escape(P.AFFIL.split(',')[0]) + r"} \\ " +
                      r"\textit{Vellore Institute of Technology}\\ Vellore, India \\ " + tex_escape(mail) + "}")
    L.append("\n\\and\n".join(blocks) + "}")
    L += [r"\maketitle", r"\begin{abstract}", tex_para(P.ABSTRACT), r"\end{abstract}",
          r"\begin{IEEEkeywords}", tex_escape(", ".join(P.INDEX_TERMS)), r"\end{IEEEkeywords}"]
    fig_no = 0
    for s in SECS:
        L.append(r"\section{" + tex_escape(s["h"]) + "}")
        for para in s["p"]:
            L.append(tex_para(para) + "\n")
        figs = ([s["fig"]] if "fig" in s else []) + s.get("figs", [])
        for cap, hdr, rows in s.get("tables", []):
            L += [r"\begin{table}[!t]", r"\caption{" + tex_escape(cap) + "}", r"\centering", r"\footnotesize",
                  r"\begin{tabular}{l" + "r" * (len(hdr) - 1) + "}", r"\toprule",
                  " & ".join(r"\textbf{" + tex_escape(h) + "}" for h in hdr) + r" \\", r"\midrule"]
            L += [" & ".join(tex_escape(str(c)) for c in row) + r" \\" for row in rows]
            L += [r"\bottomrule", r"\end{tabular}", r"\end{table}"]
        for fn, cap, wide in figs:
            fig_no += 1
            env = "figure*" if wide else "figure"
            width = r"0.85\textwidth" if wide else r"\columnwidth"
            L += [r"\begin{" + env + "}[!t]", r"\centering", r"\includegraphics[width=" + width + "]{" + fn + "}",
                  r"\caption{" + tex_escape(cap) + "}", r"\label{fig:" + str(fig_no) + "}", r"\end{" + env + "}"]
    L += [r"\section*{Acknowledgment}", tex_escape(P.ACK)]
    L += [r"\begin{thebibliography}{00}"]
    L += [r"\bibitem{" + k + "} " + bib_entry(REFS[k]) for k in ORDER]
    L += [r"\end{thebibliography}", r"\end{document}"]
    return "\n".join(L)


def compile_tex() -> Path:
    BUILD.mkdir(parents=True, exist_ok=True)
    shutil.copy(CLS, BUILD / "IEEEtran.cls")
    figs = {fn for s in SECS for fn, *_ in ([s["fig"]] if "fig" in s else []) + s.get("figs", [])}
    for fn in figs:
        shutil.copy(FIG / fn, BUILD / fn)
    (BUILD / "main.tex").write_text(build_tex(), encoding="utf-8")
    subprocess.run([str(TECTONIC), "-X", "compile", "--keep-logs", "main.tex"], cwd=BUILD, check=True)
    pdf = OUT / f"{PREFIX}_IEEE_Paper.pdf"
    shutil.copy(BUILD / "main.pdf", pdf)
    with zipfile.ZipFile(OUT / f"{PREFIX}_IEEE_Paper_LaTeX.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for p in BUILD.iterdir():
            if p.suffix in {".tex", ".cls", ".png"}:
                z.write(p, p.name)
        z.writestr("README.txt", "Compile with: pdflatex main.tex (twice) or tectonic main.tex. Class: IEEEtran v1.8b (CTAN).\n")
    return pdf


# ------------------------------------------------------------------ Word copy
def docx_runs(par, text, size=10):
    text = re.sub(r"\s*\[\[cite:([^\]]+)\]\]", lambda m: " [" + "], [".join(str(NUM[k]) for k in m.group(1).split(",")) + "]", text)
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = par.add_run(part); r.bold = i % 2 == 1; r.font.size = Pt(size); r.font.name = "Times New Roman"


def set_cols(section, n):
    cols = section._sectPr.xpath("./w:cols")
    el = cols[0] if cols else OxmlElement("w:cols")
    el.set(qn("w:num"), str(n)); el.set(qn("w:space"), "360")
    if not cols:
        section._sectPr.append(el)


def build_docx() -> Path:
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(4.3)
    sec.left_margin = sec.right_margin = Cm(1.43)
    cp = d.core_properties
    cp.title, cp.author, cp.subject = P.TITLE, "Sharvesh Prabhu; Nithin S", "IEEE conference paper - BACSE301 project outcome"
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(P.TITLE); r.font.size = Pt(24); r.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(12)
    t = d.add_table(rows=1, cols=len(P.AUTHORS))
    for i, (name, reg, mail) in enumerate(P.AUTHORS):
        c = t.rows[0].cells[i].paragraphs[0]; c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for line, it in [(f"{name} ({reg})", False), ("School of Computer Science and Engineering", True),
                         ("Vellore Institute of Technology", True), ("Vellore, India", False), (mail, False)]:
            rr = c.add_run(line + "\n"); rr.italic = it; rr.font.size = Pt(11 if line.startswith(name) else 10)
    new = d.add_section(WD_SECTION.CONTINUOUS); set_cols(new, 2)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract\u2014"); r.bold = True; r.italic = True; r.font.size = Pt(9)
    docx_runs(p, P.ABSTRACT, 9)
    for run in p.runs[1:]:
        run.bold = True
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Index Terms\u2014"); r.bold = True; r.italic = True; r.font.size = Pt(9)
    r = p.add_run(", ".join(P.INDEX_TERMS)); r.bold = True; r.font.size = Pt(9)
    fig_no = 0
    tab_no = 0
    for si, s in enumerate(SECS, 1):
        h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_before = Pt(8); h.paragraph_format.space_after = Pt(4)
        r = h.add_run(f"{roman(si)}. {s['h'].upper()}"); r.font.size = Pt(10)
        for para in s["p"]:
            q = d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            q.paragraph_format.first_line_indent = Cm(0.36)
            docx_runs(q, para)
        for cap, hdr, rows in s.get("tables", []):
            tab_no += 1
            c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_before = Pt(6)
            rr = c.add_run(f"TABLE {roman(tab_no)}\n{cap.upper()}"); rr.font.size = Pt(8)
            tb = d.add_table(rows=1, cols=len(hdr)); set_cell_borders(tb)
            for i, hh in enumerate(hdr):
                x = tb.rows[0].cells[i].paragraphs[0].add_run(hh); x.bold = True; x.font.size = Pt(8)
            for row in rows:
                cells = tb.add_row().cells
                for i, v in enumerate(row):
                    x = cells[i].paragraphs[0].add_run(str(v)); x.font.size = Pt(8)
        figs = ([s["fig"]] if "fig" in s else []) + s.get("figs", [])
        for fn, cap, wide in figs:
            fig_no += 1
            q = d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_before = Pt(6)
            q.add_run().add_picture(str(FIG / fn), width=Cm(8.4))
            c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rr = c.add_run(f"Fig. {fig_no}. {cap}"); rr.font.size = Pt(8)
    h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_before = Pt(8)
    h.add_run("ACKNOWLEDGMENT").font.size = Pt(10)
    q = d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; docx_runs(q, P.ACK)
    h = d.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_before = Pt(8)
    h.add_run("REFERENCES").font.size = Pt(10)
    for k in ORDER:
        p_ = REFS[k]
        q = d.add_paragraph(); q.paragraph_format.left_indent = Cm(0.6); q.paragraph_format.first_line_indent = Cm(-0.6)
        txt = re.sub(r"\\url\{([^}]*)\}", r"\1", bib_entry(p_)).replace("``", "\u201c").replace("''", "\u201d")
        txt = txt.replace(r"\&", "&").replace(r"\_", "_").replace(r"\%", "%")
        rr = q.add_run(f"[{NUM[k]}] {txt}"); rr.font.size = Pt(8)
    out = OUT / f"{PREFIX}_IEEE_Paper.docx"
    d.save(out)
    return out


if __name__ == "__main__":
    pdf = compile_tex()
    dx = build_docx()
    import pymupdf
    print("pages:", pymupdf.open(pdf).page_count, "refs:", len(ORDER), pdf.name, dx.name)
