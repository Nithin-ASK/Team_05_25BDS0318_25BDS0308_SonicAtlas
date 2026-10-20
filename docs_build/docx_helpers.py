"""python-docx helpers that reproduce the faculty Sample Report layout."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = "Times New Roman"


def _roman(n: int) -> str:
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals * 4:
        while n >= v:
            out += s
            n -= v
    return out


def roman(n: int) -> str:
    res, table = "", [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
                      (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    for v, s in table:
        while n >= v:
            res += s
            n -= v
    return res


def set_cell_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def set_col_widths(table, widths_cm) -> None:
    """Fixed layout so Word respects the column widths."""
    table.autofit = False
    k = min(1.0, 15.9 / sum(widths_cm))
    widths_cm = [w * k for w in widths_cm]
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true")
    trPr.append(el)


def add_field(run, instr: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)


def add_hyperlink(paragraph, url: str, text: str, size: float | None = None):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); rPr.append(fonts)
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0000EE"); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); paragraph._p.append(h)
    return h


class Report:
    """Thin wrapper keeping figure/table counters and the layout rules in one place."""

    def __init__(self, title: str, author: str, subject: str):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.headings: list[tuple[str, str]] = []  # (level, text)
        cp = self.doc.core_properties
        cp.title, cp.author, cp.subject = title, author, subject
        cp.keywords = "EDA, Spotify, audio features, PCA, t-SNE, LLE, DuckDB-WASM"
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        st = self.doc.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5; pf.space_after = Pt(10); pf.space_before = Pt(0)
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
            h = self.doc.styles[name]
            h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.element.rPr.rFonts.set(qn("w:asciiTheme"), FONT) if False else None
            rf = h.element.rPr.rFonts
            for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(att), FONT)
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(att)) is not None:
                    del rf.attrib[qn(att)]
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15
        # page number bottom right
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run(); r.font.size = Pt(10); r.font.name = FONT
        add_field(r, "PAGE")

    # ---- text -------------------------------------------------------------
    def para(self, text: str = "", bold_label: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, italic=False, bold=False, keep_next=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        if bold_label:
            r = p.add_run(bold_label); r.bold = True
        self._runs(p, text, size=size, italic=italic, bold=bold)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    def _runs(self, p, text, size=None, italic=False, bold=False):
        """Support **bold** spans inside text."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run(part)
            r.bold = bold or (i % 2 == 1)
            r.italic = italic
            if size:
                r.font.size = Pt(size)

    def bullet(self, label: str, text: str):
        p = self.doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        if label:
            r = p.add_run(label + ": "); r.bold = True
        self._runs(p, text)
        return p

    def numbered(self, items: list[str]):
        for i, t in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(4)
            self._runs(p, f"{i}. {t}")

    # ---- headings -----------------------------------------------------------
    def h1(self, text: str, new_page: bool = True):
        if new_page:
            self.page_break()
        p = self.doc.add_heading(text.upper(), level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.headings.append(("1", text.upper()))
        return p

    def h2(self, text: str, center: bool = False):
        p = self.doc.add_heading(text, level=2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        self.headings.append(("2", text))
        return p

    def h3(self, text: str):
        p = self.doc.add_heading(text, level=3)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return p

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ---- figures and tables ----------------------------------------------------
    def figure(self, path: Path, caption: str, width_in: float = 6.2) -> int:
        self.fig_n += 1
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Inches(width_in))
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.line_spacing = 1.0
        r = c.add_run(f"Fig. {self.fig_n}. {caption}"); r.font.size = Pt(10)
        return self.fig_n

    def next_fig(self) -> int:
        return self.fig_n + 1

    def next_tab(self) -> int:
        return self.tab_n + 1

    def table(self, header: list[str], rows: list[list], caption: str, widths_cm: list[float] | None = None,
              size: float = 10, links: dict[tuple[int, int], str] | None = None, bold_first_col=False) -> str:
        self.tab_n += 1
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_borders(t)
        repeat_header(t.rows[0])
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(h)); r.bold = True; r.font.size = Pt(size)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, v in enumerate(row):
                p = cells[ci].paragraphs[0]
                p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if isinstance(v, str) and len(v) > 14 else WD_ALIGN_PARAGRAPH.CENTER
                url = (links or {}).get((ri, ci))
                if url:
                    r = p.add_run(str(v) + " "); r.font.size = Pt(size)
                    add_hyperlink(p, url, url, size=size - 1)
                else:
                    self._runs(p, str(v), size=size, bold=bold_first_col and ci == 0)
        if widths_cm:
            set_col_widths(t, widths_cm)
        if len(rows) <= 16:  # keep short tables on one page
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
            for row in t.rows:
                no_split(row)
        c = self.doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before = Pt(4); c.paragraph_format.line_spacing = 1.0
        r = c.add_run(f"TABLE {roman(self.tab_n)}: {caption}"); r.font.size = Pt(10)
        return roman(self.tab_n)

    def save(self, path: Path):
        self.doc.save(str(path))


def to_pdf_word(docx: Path, pdf: Path) -> None:
    """Export with Microsoft Word through COM (LibreOffice is not installed)."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx.resolve()), ReadOnly=True)
        d.ExportAsFixedFormat(str(pdf.resolve()), 17)
        d.Close(False)
    finally:
        word.Quit()


def pptx_to_pdf(pptx: Path, pdf: Path) -> None:
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx.resolve()), WithWindow=False)
        pres.SaveAs(str(pdf.resolve()), 32)
        pres.Close()
    finally:
        app.Quit()
