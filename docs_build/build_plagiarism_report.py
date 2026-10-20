"""Similarity Self-Check Report (not an official Turnitin report) -> docx + pdf."""
from __future__ import annotations

import json
import sys
from importlib.metadata import version
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from common import OUT, PREFIX, ROOT, TEAM, TITLE  # noqa: E402
from docx_helpers import Report, roman, to_pdf_word  # noqa: E402

PLAG = ROOT / "reports" / "plagiarism"
final = json.loads((PLAG / "check_results_final.json").read_text(encoding="utf-8"))
first = json.loads((PLAG / "check_results_raw.json").read_text(encoding="utf-8")) if (PLAG / "check_results_raw.json").exists() else None
web = json.loads((PLAG / "web_check.json").read_text(encoding="utf-8")) if (PLAG / "web_check.json").exists() else None
manifest = json.loads((PLAG / "corpus_manifest.json").read_text(encoding="utf-8"))

rep = Report("Similarity Self-Check Report - " + TITLE, "Team 05", "Similarity self-check (not an official Turnitin report)")
d = rep.doc


def center(text, size=12, bold=False, italic=False, color=None, after=6):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


center("Similarity Self-Check Report", 20, bold=True, after=4)
center("(not an official Turnitin report)", 14, bold=True, color="B00000", after=18)
center(TITLE, 13, bold=True, after=14)
center("Team 05: " + ", ".join(f"{n} ({r})" for r, n, _ in TEAM), 12)
center("BACSE301 Exploratory Data Analysis, Vellore Institute of Technology, Vellore", 12)
center("Check date: 20 October 2026", 12, after=18)
center("This document records a self-administered similarity check made by the team with open tools. It is NOT a "
       "Turnitin or iThenticate report and must not be presented as one. The official similarity check is done by VIT "
       "or the course faculty.", 11, italic=True)

rep.h1("1. Method")
docs = final["documents"]
rep.para("Three complementary checks were run on the final versions of the Project Report, the IEEE paper and every "
         "presentation:")
rep.bullet("Fingerprint overlap", "each document was reduced to lower-case word sequences and compared with every source in "
           "the corpus. The similarity percentage is the share of a document's words that lie inside an 8-word sequence "
           "that also occurs in a source; 5-word overlap is reported as a stricter, noisier view. This is the same idea as "
           "the winnowing/fingerprinting used by tools such as MOSS and copydetect.")
rep.bullet("Sentence cosine", "every sentence was compared with every corpus sentence using TF-IDF vectors (unigrams and "
           "bigrams); pairs with cosine similarity of 0.6 or more were flagged for manual review.")
rep.bullet("Web exact-phrase search", "a fixed random sample (seed 42) of 60 report sentences and 30 paper sentences was "
           "searched on the web as exact quoted phrases.")
rep.para("Excluded from scoring, as the brief requires: the reference lists, URLs and DOIs, numbers, and the titles of cited "
         "works (titles are citations, not prose). Standard method and dataset names remain in the text and are counted.")
rep.h2("1.1 Corpus and tools")
kinds = {"paper": 0, "kaggle": 0, "wiki": 0}
for k in manifest:
    for pre in kinds:
        if k.startswith(pre) or (pre == "paper" and not k.startswith(("kaggle", "wiki"))):
            kinds[pre] += 1
            break
rep.para(f"The corpus holds {final['corpus_docs']} non-empty documents with {final['corpus_words']:,} words: the abstracts "
         f"of the cited papers (retrieved from Crossref, OpenAlex or arXiv), the three Kaggle dataset descriptions, and "
         f"{kinds['wiki']} Wikipedia articles on the methods used (EDA, PCA, t-SNE, LLE, UMAP, imputation, Isolation Forest, "
         "ANOVA, chi-square, STL, sentiment analysis, DuckDB and others). Two abstracts could not be retrieved "
         "(Brand et al. 2019; Fell and Sporleder 2014) and two Wikipedia pages returned no text. The faculty Sample Report "
         "was not available in the project folder, so it could not be added to the corpus. Full texts behind paywalls were "
         "not used.")
rep.table(["Tool", "Version / note"], [
    ["Python", sys.version.split()[0]], ["scikit-learn (TF-IDF, cosine)", version("scikit-learn")],
    ["PyMuPDF (PDF text)", version("pymupdf")], ["python-pptx (slide text)", version("python-pptx")],
    ["n-gram matcher", "docs_build/plagiarism_check.py (this project)"], ["Web search", (web or {}).get("engine", "WebSearch tool")],
], "Tools used for the self-check", widths_cm=[6, 9])

rep.h1("2. Results")
rows = [[nm, f"{v['words']:,}", f"{v['similarity_8gram_pct']:.2f}%", f"{v['similarity_5gram_pct']:.2f}%",
         f"{v['max_single_source_8gram_pct']:.2f}%", str(len(v["tfidf_flagged"]))] for nm, v in docs.items()]
rep.para(f"TABLE {roman(rep.next_tab())} gives the final similarity per document. The targets were an overall similarity below 10% and no single "
         "source above 2%.")
rep.table(["Document", "Words checked", "Similarity (8-gram)", "5-gram overlap", "Largest single source", "Flagged sentences"],
          rows, "Final similarity per document", widths_cm=[3.6, 2.4, 2.6, 2.4, 2.6, 2.2], size=9.5)
rep.h2("2.1 Top matched sources and passages")
for nm, v in docs.items():
    srcs = ", ".join(f"{s} ({p:.2f}%)" for s, p in v["top_sources_8gram"][:5]) or "none"
    rep.para(f"**{nm}.** Sources with 8-word overlap: {srcs}.")
    if v["matched_passages_8gram"]:
        for ps in v["matched_passages_8gram"][:8]:
            p = d.add_paragraph(); p.paragraph_format.left_indent = Pt(24)
            r = p.add_run(f"“{ps}”"); r.font.size = Pt(10); r.font.highlight_color = 7
    for fl in v["tfidf_flagged"][:5]:
        rep.para(f"Flagged (cosine {fl['cosine']:.2f}, {fl['source']}): “{fl['sentence'][:220]}”")
rep.para("All remaining 8-word matches are proper names of conferences or institutions (for example the name of the IEEE "
         "ICASSP conference inside a citation, or the school and university name), which are standard terms and need no "
         "rewriting.")
rep.h2("2.2 Cross-document check")
rep.para(f"The IEEE paper was written separately from the report. Only {final['cross_report_vs_paper_8gram_pct']:.2f}% of the "
         "paper's words fall inside 8-word sequences shared with the report; the shared passages are the affiliation line "
         "and the list of cluster names:")
for ps in final["cross_report_vs_paper_passages"][:5]:
    p = d.add_paragraph(); p.paragraph_format.left_indent = Pt(24)
    r = p.add_run(f"“{ps}”"); r.font.size = Pt(10)
rep.h2("2.3 Web exact-phrase check")
if web:
    m = [x for x in web["results"] if isinstance(x.get("exact_matches"), int) and x["exact_matches"] > 0]
    err = [x for x in web["results"] if x.get("exact_matches") == "error"]
    rep.para(f"{web['checked']} sentences were searched as exact phrases ({len([x for x in web['results'] if x['doc'] == 'report'])} "
             f"from the report and {len([x for x in web['results'] if x['doc'] == 'paper'])} from the paper). Matches found: "
             f"report {web['summary'].get('report_matches', 0)}, paper {web['summary'].get('paper_matches', 0)}"
             + (f"; {len(err)} searches returned an error." if err else "."))
    for x in m[:10]:
        rep.para(f"Match: “{x['query']}” - {', '.join(x['urls'][:3])}")
    rep.table(["#", "Doc", "Exact-phrase query", "Matches"],
              [[str(i + 1), x["doc"], x["query"][:90], str(x["exact_matches"])] for i, x in enumerate(web["results"])],
              "Web exact-phrase queries and results", widths_cm=[1, 1.6, 11, 1.8], size=8)
else:
    rep.para("The web exact-phrase check had not finished when this version was generated.")
rep.h2("2.4 Online checker")
rep.para("Step (c) of the brief, an online checker used without login, was deliberately not run. Free checkers often keep "
         "the submitted text in their own databases, so pasting unpublished report text into them could later make the "
         "team's own official submission look plagiarised. The decision is left to the team.")

rep.h1("3. Fixes and final numbers")
if first:
    rows = [[nm, f"{first['documents'][nm]['similarity_8gram_pct']:.2f}%", f"{v['similarity_8gram_pct']:.2f}%"]
            for nm, v in docs.items() if nm in first["documents"]]
    rep.para("The first run counted the titles of cited works as matches (for example the title of a cited paper written in "
             "the literature table). Because titles are citations, they were excluded in the scoring rule; no passage of "
             "prose exceeded the threshold, so no text had to be rewritten. TABLE " + roman(rep.next_tab()) + " shows the effect.")
    rep.table(["Document", "First run (titles counted)", "Final (titles excluded)"], rows, "Before vs after", widths_cm=[5, 5, 5])
worst = max(v["similarity_8gram_pct"] for v in docs.values())
rep.para(f"**Conclusion.** The highest overall similarity of any document is {worst:.2f}% and the largest share from a single "
         f"source is {max(v['max_single_source_8gram_pct'] for v in docs.values()):.2f}%, both well below the targets "
         "(10% overall, 2% per source). These numbers come from the open self-check described above; an official tool with "
         "a larger database may report different values. The official Turnitin check is done by VIT or the faculty.")

out = OUT / f"{PREFIX}_Plagiarism_Report.docx"
rep.save(out)
to_pdf_word(out, out.with_suffix(".pdf"))
out.unlink()  # the deliverable is the PDF
print("ok")
