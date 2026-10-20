"""One-page Submission Checklist PDF."""
import sys
from pathlib import Path

from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from common import OUT, PREFIX  # noqa: E402
from docx_helpers import Report, to_pdf_word  # noqa: E402

rep = Report("Submission Checklist - Team 05", "Team 05", "Submission checklist")
d = rep.doc
sec = d.sections[0]
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.4))
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"{PREFIX} - Submission Checklist"); r.bold = True; r.font.size = Pt(13)
p.paragraph_format.space_after = Pt(2)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SonicAtlas, BACSE301 Exploratory Data Analysis, VIT Vellore, October 2026"); r.font.size = Pt(9)
p.paragraph_format.space_after = Pt(4)

F = lambda s: f"{PREFIX}_{s}"
rows = [
    ["Project_Report.pdf / .docx", "Case study report (Lab Assessment-7)", "Faculty, by 20.10.2026", "[ ]"],
    ["Review2_PPT.pptx / .pdf", "Review-2 presentation, 10 slides", "E-mail to faculty (draft 1)", "[ ]"],
    ["Review3_PPT.pptx / .pdf", "Review-3 presentation, 10 slides", "E-mail to faculty (draft 2)", "[ ]"],
    ["Final_PPT.pptx / .pdf", "Final viva presentation, 10 slides", "Final submission (draft 3)", "[ ]"],
    ["Review1_PPT.pptx / .pdf", "Bonus: Review-1 structure, 8 slides", "If the faculty asks", "[ ]"],
    ["IEEE_Paper.pdf / .docx / _LaTeX.zip", "Project outcome paper (prepared for submission)", "Conference, after you choose one", "[ ]"],
    ["Outcome_Proof.pdf", "Evidence pack for the outcome", "Faculty; add pending proofs later", "[ ]"],
    ["Plagiarism_Report.pdf", "Similarity SELF-check (not Turnitin)", "Faculty (optional)", "[ ]"],
    ["Source_Code.zip", "Clean code snapshot, no data or secrets", "Faculty / GitHub", "[ ]"],
    ["Demo_Video.mp4", "Website walkthrough with captions", "Faculty, link in report", "[ ]"],
    ["Email_Drafts.txt", "Three ready-to-send e-mails (not sent)", "You send them", "[ ]"],
    ["MANUAL_STEPS_FOR_ME.pdf / .txt", "Step-by-step guide to your own jobs", "Read first", "[ ]"],
    ["Viva_Questions.pdf", "20 likely viva questions with short answers", "Practise", "[ ]"],
]
rep.table(["File (after " + PREFIX[:13] + "...)", "What it is", "Where / when to submit", "Done"], rows,
          "Deliverables in the deliverables folder", widths_cm=[5.6, 6.0, 4.8, 1.2], size=8.5)
d.paragraphs[-1].paragraph_format.space_after = Pt(2)
p = d.add_paragraph(); r = p.add_run("Things only you can do"); r.bold = True; r.font.size = Pt(10)
p.paragraph_format.space_after = Pt(1)
for t in [
    "Confirm the school name on the cover (we used SCHOOL OF COMPUTER SCIENCE AND ENGINEERING) and the course name; write the faculty name where it shows ____.",
    "Paste the official VIT logo into the empty box on the report cover (we cannot legally copy it).",
    "Replace [GitHub Link], [Live Website Link], [Demo Video Link] and [Paper Link] in the report, decks and e-mails; put Sharvesh's e-mail where it shows <email>.",
    "Push the repository to GitHub and enable Pages (Settings -> Pages -> Source: GitHub Actions); upload tracks_full_web.parquet as a Release asset if you want the full data tier.",
    "Send the three e-mails yourself; check every attachment name and link first.",
    "Pick a venue (IEEE CONECCT 2027: deadline 15 Dec 2026) and run IEEE PDF eXpress with that conference ID; submit the paper yourself.",
    "The official Turnitin check is done by VIT or the faculty. Our report is a self-check only and says so on its cover.",
    "Add to the Outcome Proof only real documents: submission receipt, acceptance e-mail, certificate, Turnitin report.",
]:
    q = d.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = 1.0
    rr = q.add_run(t); rr.font.size = Pt(8.5)
p = d.add_paragraph(); r = p.add_run("Run the website locally: cd web && npm ci && npm run dev  (http://localhost:5173). "
                                      "Rebuild documents: see Makefile (make report ppt paper proof plagcheck)."); r.font.size = Pt(8.5)
out = OUT / F("Submission_Checklist.docx")
rep.save(out)
to_pdf_word(out, out.with_suffix(".pdf"))
out.unlink()
import pymupdf
print("checklist pages:", pymupdf.open(out.with_suffix(".pdf")).page_count)
