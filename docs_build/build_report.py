"""Generate the Project Report (.docx + .pdf) from results.json and literature.json.

Two passes: build -> export PDF -> read real page numbers of each heading -> rebuild the TOC -> export again.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import fitz  # PyMuPDF
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from common import (COURSE, COURSE_CODE, FIG, MONTH, OUT, PREFIX, ROOT, SCHOOL, TEAM, TITLE, UNIV, LIT, R, f, links, n, pct, pval)  # noqa: E402
from docx_helpers import Report, add_hyperlink, roman, set_cell_borders, set_col_widths, to_pdf_word  # noqa: E402

r = R()
lit = LIT()
L = links()
A, C, D, M, DR, TS, T = (r["audit"], r["cleaning"], r["descriptive"], r["multivariate"], r["dimred"], r["timeseries"], r["tests"])
LY, PR, WEB, NET = r.get("lyrics"), r.get("profiling"), r.get("web"), r.get("network")

# ---------------------------------------------------------------- references
papers = lit["papers"]
software = lit["software"]
REFS: list[dict] = papers + software
REF_NUM = {p["id"]: i + 1 for i, p in enumerate(REFS)}


def cite(*ids: str) -> str:
    nums = sorted(REF_NUM[i] for i in ids if i in REF_NUM)
    return "[" + "], [".join(map(str, nums)) + "]" if nums else ""


def cite_topic(t: str, k: int = 3) -> str:
    return cite(*[p["id"] for p in papers if p.get("topic") == t][:k])


def cite_sw(word: str) -> str:
    return cite(*[s["id"] for s in software if word.lower() in (s.get("title", "") + " " + s["id"]).lower()][:1])


def ref_text(p: dict) -> str:
    venue = p.get("venue") or ""
    s = f"{p.get('authors', '')}, “{p.get('title', '')},” {venue}"
    if p.get("doi") and p["doi"] not in venue:
        s += f", doi: {p['doi']}"
    return s.rstrip(".") + "."


rep = Report(TITLE, "Sharvesh Prabhu (25BDS0318); Nithin S (25BDS0308)", f"{COURSE_CODE} {COURSE} - Project Report, Team 05")
doc = rep.doc
FOCUS = M["focus_genres"]

# ================================================================= COVER


def cover():
    def line(text, size=12, bold=False, italic=False, before=0, after=6):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text); run.bold = bold; run.italic = italic; run.font.size = Pt(size)
        return p
    line(TITLE, 16, bold=True, before=24, after=30)
    line("Submitted by", 13, bold=True, italic=True, after=10)
    for reg, name, _ in TEAM:
        line(f"{reg}    {name}", 13, after=4)
    line("In partial fulfillment of the requirements for the completion of the", 12, bold=True, italic=True, before=30, after=8)
    line("Project-based Course", 13, bold=True, after=8)
    line(f"{COURSE_CODE} {COURSE}", 13, bold=True, after=24)
    # logo placeholder box (the official logo must be pasted by the team)
    t = doc.add_table(rows=1, cols=1); set_cell_borders(t)
    t.alignment = 1
    cell = t.rows[0].cells[0]; cell.width = Cm(6)
    t.rows[0].height = Cm(3.2)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("[Paste official VIT logo here]"); run.italic = True; run.font.size = Pt(10)
    line(SCHOOL, 13, bold=True, before=26, after=6)
    line(UNIV, 15, bold=True, after=20)
    line(MONTH, 13, bold=True)


# ================================================================= TOC


TOC_ITEMS = [
    ("1.", "ABSTRACT", True), ("2.", "INTRODUCTION", True), ("", "2.1 Background", False), ("", "2.2 Motivations", False),
    ("", "2.3 Scope of the Project", False), ("3.", "PROJECT DESCRIPTION AND GOALS", True), ("", "3.1 Literature Review", False),
    ("", "3.2 Gaps Identified", False), ("", "3.3 Objectives", False), ("", "3.4 Problem Statement", False),
    ("4.", "PROPOSED WORK", True), ("", "4.1 Data Acquisition, Cleaning and Preprocessing of Large-Scale Spotify Audio and Lyrics Data", False),
    ("", "4.2 Descriptive Statistics, Multivariate Visual Storytelling, Dimensionality Reduction and Time-Series Dynamics of Musical Characteristics", False),
    ("", "4.3 The SonicAtlas Interactive Query Platform: Architecture, Text Mining, Hypothesis Testing and Evaluation Framework", False),
    ("5.", "CONCLUSION", True), ("6.", "PROJECT OUTCOME", True), ("7.", "REFERENCES", True),
]


def toc(pages: dict[str, int] | None):
    rep.page_break()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TABLE OF CONTENTS"); run.bold = True; run.font.size = Pt(14)
    t = doc.add_table(rows=1, cols=3); set_cell_borders(t); t.alignment = 1
    for i, h in enumerate(["Sl.No", "Contents", "Page No."]):
        c = t.rows[0].cells[i]; pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(h); rr.bold = True
    for num, text, main in TOC_ITEMS:
        cells = t.add_row().cells
        key = text.split(" ", 1)[0] if not main else text
        pg = (pages or {}).get(key, "")
        for i, v in enumerate([num, text, str(pg)]):
            pp = cells[i].paragraphs[0]; pp.paragraph_format.space_after = Pt(2); pp.paragraph_format.line_spacing = 1.15
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            rr = pp.add_run(v); rr.bold = main and i < 2
    set_col_widths(t, [1.8, 12.2, 2.2])


# ================================================================= BODY


def abstract():
    rep.h1("1. Abstract")
    rep.para(
        f"Streaming platforms describe every track with a vector of model-derived audio descriptors, yet most public "
        f"analyses of these descriptors stay inside static notebooks that only their authors can explore. This project "
        f"presents SonicAtlas, an exploratory data analysis (EDA) study and an in-browser query platform built on "
        f"{n(A['a_rows'])} Spotify tracks released between {A['a_year_min']} and {A['a_year_max']}, a genre-labelled "
        f"table of {n(C['rows_genre_tracks'])} unique tracks across {A['g_genres']} genres, and a sample of Genius lyrics. "
        f"Impossible values were repaired with MICE imputation, skewed fields were transformed, and outliers were flagged "
        f"rather than removed. Descriptive statistics follow the course conventions (n-1 variance, Tukey quartiles, G1 and "
        f"G2). Energy is tightly tied to loudness (r = {f(M['pairs']['loudness~energy']['pearson'], 2)}) and opposed to "
        f"acousticness (r = {f(M['pairs']['energy~acousticness']['pearson'], 2)}). Between {TS['year_window'][0]} and "
        f"{TS['year_window'][1]} mean loudness rose from {f(TS['first_last']['loudness'][0], 1)} dB to "
        f"{f(TS['first_last']['loudness'][1], 1)} dB while acousticness fell from {f(TS['first_last']['acousticness'][0], 2)} "
        f"to {f(TS['first_last']['acousticness'][1], 2)}. PCA needs {DR['n90']} components for 90% of the variance; t-SNE "
        f"keeps local structure best (trustworthiness {f(DR['comparison']['t-SNE']['trustworthiness_k10'], 2)}). Genre "
        f"explains {pct(100 * T['anova']['eta_squared'])} of the variance in danceability (ANOVA, {pval(T['anova']['p'])}). "
        f"All results are served through a static website that runs SQL on Parquet files inside the browser with DuckDB-WASM.")
    rep.para("**Keywords:** exploratory data analysis, Spotify audio features, dimensionality reduction, time series, "
             "lyrics mining, DuckDB-WASM, interactive visual analytics")


def introduction():
    rep.h1("2. Introduction")
    rep.h2("2.1 Background")
    rep.para(
        "Digital music services now hold catalogues of tens of millions of recordings. To organise and recommend this "
        "material, Spotify and similar services compute numerical descriptors for each recording: danceability, energy, "
        "valence (musical positivity), acousticness, instrumentalness, speechiness, liveness, loudness in decibels, tempo "
        "in beats per minute, key, mode and duration. Until late 2024 these values were available through a public Web API, "
        "and researchers and hobbyists collected them into large open tables that are still shared on Kaggle. Because the "
        "audio-features endpoint is no longer open to new applications, these archived tables have become the main public "
        "window onto the acoustic profile of recorded music.")
    rep.para(
        f"Musicologists and data scientists have used such descriptors to ask how popular music has changed over time "
        f"{cite_topic('d', 2)}, whether success can be predicted from sound {cite_topic('c', 2)}, and how genres differ "
        f"{cite_topic('b', 2)}. Lyrics add a second, textual view of the same songs; sentiment tools such as VADER "
        f"{cite_topic('e', 1)} make it possible to compare the emotion of the words with the emotion of the sound. "
        f"Exploratory Data Analysis, the practice of summarising and visualising data before building models, is the natural "
        f"framework for this kind of open-ended question, and modern dimensionality-reduction methods such as PCA, t-SNE and "
        f"LLE {cite_topic('g', 3)} allow a ten-dimensional audio space to be drawn as a readable map.")
    rep.h2("2.2 Motivations")
    rep.para(
        f"Three observations motivated this project. First, most published analyses of Spotify features are static: a "
        "paper or notebook shows a fixed set of charts, and a reader who wants a different genre, decade or feature has to "
        "download the data and repeat the work. Second, automated EDA tools such as YData-Profiling, Sweetviz and D-Tale "
        f"{cite('ydataprofiling', 'sweetviz', 'dtale', 'handout_autoeda')} produce useful overviews but know nothing about music, so they cannot answer domain questions such as “when did "
        "recorded music become louder?”. Third, a browser-based analytical database (DuckDB compiled to WebAssembly) now "
        "makes it possible to run SQL over hundreds of thousands of rows on the reader’s own laptop with no server, which "
        "removes the hosting cost that usually stops student projects from becoming shared tools.")
    rep.para(
        "The course itself provided a further motivation. The BACSE301 syllabus covers descriptive statistics, correlation, "
        "multivariate plots (parallel coordinates and Andrews curves), PCA, SVD, t-SNE, LLE and time-series decomposition. "
        "A music catalogue exercises every one of these techniques on real, messy, large-scale data, and the faculty’s "
        "reading on a Streamlit-based EDA web application " + cite("dasilva2025parkinson") +
        " showed that packaging an analysis as an interactive application is itself a recognised research contribution.")
    rep.h2("2.3 Scope of the Project")
    rep.para(
        "The project covers the full path from public raw files to a published, queryable tool. Its scope is organised into "
        "three core areas: (i) data engineering, which includes acquisition of three public Kaggle datasets, auditing, type "
        "conversion, MICE imputation of impossible values, transformation, scaling, outlier flagging and joining genre labels "
        "to the main catalogue; (ii) statistical and visual exploration, which includes descriptive statistics in the course "
        "conventions, correlation analysis, multivariate genre fingerprints, PCA/SVD, t-SNE, LLE and UMAP maps, clustering, "
        "time-series trend, seasonality and change-point analysis, hypothesis tests, feature selection and lyrics text mining; "
        "and (iii) the SonicAtlas platform, a static web application in which users filter, query, compare and map the data "
        "inside their own browser.")
    rep.para(
        "The project does not collect new audio, does not call the restricted Spotify Web API, and does not train a "
        "production recommender. Spotify’s descriptors are themselves outputs of proprietary models, so all conclusions are "
        "about these descriptors rather than about the underlying sound waves.")


def literature():
    rep.h1("3. Project Description and Goals")
    rep.h2("3.1 Literature Review")
    rep.para(
        f"The review below covers {len(papers)} published works that were opened and checked through their DOI, arXiv or "
        f"publisher pages; candidate papers were located with the research-discovery tools listed in the course handout "
        f"{cite('handout_aitools')}, and the serial number of each row in the table equals its reference number in Section 7. They span audio-feature analysis, genre classification, popularity and trend studies, the "
        f"long-term evolution of loudness and timbre, lyrics sentiment, recommendation, dimensionality reduction, automated "
        f"EDA and in-browser analytics. TABLE {roman(rep.next_tab())} summarises "
        f"each work, why it matters for SonicAtlas, and the drawback that SonicAtlas tries to address.")
    rows, links_ = [], {}
    for i, p in enumerate(papers):
        venue = re.sub(r"\s+", " ", p.get("venue", ""))
        rows.append([str(i + 1), p.get("authors", ""), p.get("title", ""), venue, p.get("why_it_matters", ""),
                     p.get("description", ""), p.get("drawbacks", "")])
        if p.get("url"):
            links_[(i, 3)] = p["url"]
    rep.table(["Sr. No", "Author's Name", "Title of Paper", "Journal Name and link", "Why this paper matters?",
               "Description of their Work", "Drawbacks Found in that"], rows,
              "Literature review of related work", widths_cm=[0.9, 2.0, 2.3, 2.5, 2.7, 3.3, 2.5], size=8, links=links_)


def gaps_objectives():
    rep.h2("3.2 Gaps Identified")
    rep.para("Reading the drawbacks column of Table I together gives five recurring gaps:")
    rep.bullet("Static presentation", "trend and popularity studies publish fixed figures; readers cannot re-slice the data by genre, decade or feature without repeating the whole pipeline.")
    rep.bullet("Sound and words studied apart", "audio-feature papers rarely look at lyrics, and lyrics papers rarely use audio descriptors, so the agreement between musical and lyrical mood is seldom measured.")
    rep.bullet("Little attention to data quality", "impossible values (zero tempo, zero time signature), re-issued duplicates and outliers are usually dropped silently or not discussed, which can bias decade averages.")
    rep.bullet("No comparison of embedding methods on music descriptors", "music maps typically use one method (often t-SNE) without reporting runtime, trustworthiness or cluster separation against PCA and LLE.")
    rep.bullet("Generic tools are not domain-aware", "automated EDA libraries and the Streamlit EDA application describe columns but cannot answer music-specific questions, and most interactive tools need a running server.")
    rep.h2("3.3 Objectives")
    rep.para("The objectives below are numbered and each has a measurable outcome:")
    rep.numbered([
        f"Build a reproducible cleaning pipeline for at least 1 million tracks that repairs impossible values with MICE (validated by masking 5% of known values and comparing RMSE with mean imputation) and flags outliers with two methods.",
        "Report per-feature descriptive statistics (mean, median, mode, n-1 variance, Tukey Q1/Q3/IQR, G1, G2, CV) for the whole catalogue, for eight representative genres and for every decade from the 1950s to the 2020s.",
        "Quantify the main correlations and visualise genre fingerprints with parallel coordinates, Andrews curves, radar and violin plots.",
        "Compare PCA, t-SNE, LLE and UMAP on a stratified 10,000-track sample using runtime, trustworthiness (k = 10) and silhouette, and name the acoustic clusters found by KMeans.",
        "Measure long-term trends with yearly means, STL decomposition, change-point detection and a hold-out ARIMA forecast whose error is compared with a naive baseline.",
        "Test genre differences formally (one-way ANOVA with Tukey HSD, Kruskal-Wallis, chi-square with Cramer's V) and compare filter and wrapper feature-selection rankings.",
        "Publish a static web platform in which SQL queries on the lite dataset return in under one second on a normal laptop, with every result reproducible from the notebooks.",
    ])
    rep.h2("3.4 Problem Statement")
    rep.para(
        "Large public tables of Spotify audio features and song lyrics exist, but they are noisy, partly duplicated and "
        "split across sources that share no common key, and the knowledge extracted from them is locked inside static "
        "papers and notebooks. Students, researchers and open-data users therefore lack a trustworthy, reproducible and "
        "interactive way to ask how the sound of music differs across genres and how it has changed across eras. This "
        "project addresses that problem by building a documented EDA pipeline over more than one million tracks, applying "
        "the statistical, multivariate, dimensionality-reduction and time-series methods of the course, and delivering the "
        "results as SonicAtlas, a serverless in-browser platform where the cleaned data can be filtered, queried, mapped and "
        "compared by anyone.")


def proposed_41():
    rep.h1("4. Proposed Work")
    rep.para(
        f"The proposed work is organised in three blocks that follow the methodology in Fig. {rep.next_fig()}: data "
        "engineering (Section 4.1), statistical and visual exploration (Section 4.2), and the interactive platform together "
        "with formal testing and evaluation (Section 4.3).")
    rep.figure(FIG / "fig_00_02_methodology_flow.png", "Overall methodology of the SonicAtlas project", 6.4)
    rep.h2("4.1 Data Acquisition, Cleaning and Preprocessing of Large-Scale Spotify Audio and Lyrics Data", center=True)
    rep.para("This block turns three public downloads into analysis-ready Parquet tables while keeping the raw files untouched.")
    rep.h3("4.1.1 Data Sources")
    tn = rep.next_tab()
    rep.para(f"Three public Kaggle datasets were used (TABLE {roman(tn)}): the Spotify 1.2M+ Songs table {cite('kaggle_spotify12m')}, "
             f"the Spotify Tracks Dataset {cite('kaggle_spotify_tracks')} and the Genius Song Lyrics dataset {cite('kaggle_genius_lyrics')}. "
             "They were downloaded from the public Kaggle endpoint on 20 October 2026, checksummed, and stored unchanged in data/raw. "
             f"All processing used pandas {cite('mckinney2010pandas')} and scikit-learn {cite('pedregosa2011sklearn')}.")
    ly_rows = n(LY["rows_used"]) + " (English sample)" if LY else "sample (see 4.3)"
    rep.table(["Dataset", "Kaggle source", "Rows", "Role in SonicAtlas"], [
        ["Spotify 1.2M+ Songs", "rodolfofigueroa/spotify-12m-songs", n(A["a_rows"]), "Main audio table with release year"],
        ["Spotify Tracks Dataset", "maharshipandya/-spotify-tracks-dataset", n(A["g_rows"]), f"Genre labels ({A['g_genres']} genres) and popularity"],
        ["Genius Song Lyrics", "carlosgdcj/genius-song-lyrics-with-language-information", ly_rows, "Lyrics text, language tag"],
    ], "Public datasets used in the project", widths_cm=[3.5, 5.8, 2.6, 4.6])
    rep.h3("4.1.2 Data Audit")
    tn = rep.next_tab()
    ck = A["a_checks"]
    rep.para(f"The audit (notebook 01) checked shape, types, missing values, duplicates and physically impossible values. "
             f"TABLE {roman(tn)} lists what was found. Text columns had almost no gaps "
             f"({A['a_missing_total']} missing names or album titles in the main table), but {n(ck['tempo_zero'])} tracks "
             f"had a tempo of 0 BPM and {n(ck['time_signature_zero'])} a time signature of 0, which cannot occur in real music.")
    rep.table(["Check", "Spotify 1.2M+", "Spotify Tracks (genre)"], [
        ["Rows x columns", f"{n(A['a_rows'])} x {A['a_cols']}", f"{n(A['g_rows'])} x {A['g_cols']}"],
        ["Missing cells", n(A["a_missing_total"]), n(A["g_missing_total"])],
        ["Duplicate track IDs", n(A["a_dup_id"]), n(A["g_dup_id"]) + " (same track in several genres)"],
        ["Duplicate title + artist", n(A["a_dup_title_artist"]), n(A["g_dup_title_artist"])],
        ["Tempo = 0 BPM", n(ck["tempo_zero"]), n(A["g_checks"]["tempo_zero"])],
        ["Time signature = 0", n(ck["time_signature_zero"]), n(A["g_checks"]["time_signature_zero"])],
        ["Loudness above 0 dB", n(ck["loudness_positive"]), "-"],
        ["Year recorded as 0", n(ck["year_zero_or_before_1900"]), "no year column"],
        ["Duration < 30 s / > 20 min", f"{n(ck['duration_under_30s'])} / {n(ck['duration_over_20min'])}", f"{A['g_checks']['duration_under_30s']} / -"],
        ["Popularity = 0", "no popularity column", n(A["g_checks"]["popularity_zero"])],
    ], "Data audit results", widths_cm=[5.0, 4.8, 6.4])
    rep.h3("4.1.3 Type Conversion and Feature Engineering")
    dp = C["date_precision"]
    rep.para("Several columns were converted to proper types and a few readable features were derived:")
    rep.bullet("Release date", f"parsed to a datetime; {n(dp.get('day', 0))} dates have day precision, {n(dp.get('year', 0))} only a year and {n(dp.get('month', 0))} a month. Year and decade columns were derived, and the 10 tracks recorded with year 0 were set to missing.")
    rep.bullet("Key and mode", "integer pitch classes 0-11 and the binary mode were combined into readable labels such as “C# minor”.")
    rep.bullet("Artists", "the stored Python-style list strings were parsed into real lists, giving a main artist and an artist count per track.")
    rep.bullet("Explicit flag and duration", "explicit was cast to boolean and duration in milliseconds to minutes.")
    rep.h3("4.1.4 Missing Values and MICE Imputation")
    mv = C["mice_validation"]
    tn = rep.next_tab(); fn = rep.next_fig()
    rep.para(f"Zero tempo and zero time signature were recoded as missing and imputed with Multiple Imputation by Chained "
             f"Equations (MICE) {cite_sw('mice')}, using scikit-learn’s IterativeImputer {cite_sw('scikit')} with the nine "
             f"other audio features as predictors. Missing text fields were filled with the constant “Unknown”. Because "
             f"the audio features themselves contain no gaps, the imputer was validated by hiding 5% of known values and "
             f"measuring the error (TABLE {roman(tn)}). MICE beat mean imputation for every column, by "
             f"{f(mv['energy']['improvement_pct'], 1)}% for energy, {f(mv['valence']['improvement_pct'], 1)}% for valence and "
             f"{f(mv['tempo']['improvement_pct'], 1)}% for tempo; tempo is hard to predict from the other features. Fig. {fn} "
             f"shows that imputation did not distort the tempo distribution.")
    rep.table(["Column", "Masked values", "RMSE (MICE)", "RMSE (mean)", "Improvement"],
              [[k, n(v["masked"]), f(v["rmse_mice"], 4), f(v["rmse_mean"], 4), pct(v["improvement_pct"])] for k, v in mv.items()],
              "MICE validation on artificially masked values", widths_cm=[3, 3, 3, 3, 3])
    rep.figure(FIG / "fig_02_01_mice_tempo_before_after.png", "Tempo distribution before and after MICE imputation")
    rep.h3("4.1.5 Re-releases and Joining the Genre Labels")
    j = C["join"]
    rep.para(f"The same song often appears on an original album, a deluxe edition and a compilation. The earliest release "
             f"of each normalised title and main artist was marked as canonical ({n(C['canonical_rows'])} rows); the other "
             f"{n(C['reissue_rows'])} re-issues are kept but excluded from time trends so that re-releases do not count twice.")
    rep.para(f"The genre table was first de-duplicated to {n(j['genre_unique_tracks'])} unique tracks (one track can carry "
             f"several genre tags). It was then joined to the main table by Spotify track ID ({n(j['id_matches'])} matches) "
             f"and, for the rest, by normalised title and main artist ({n(j['title_artist_matches'])} more). The overall match "
             f"rate is only {pct(j['match_rate_pct'])}: the two datasets were collected at different times from different "
             f"seed playlists. For this reason genre analyses use the genre table directly, and only "
             f"{n(C['genre_tracks_with_year'])} genre-labelled tracks also have a release year.")
    rep.h3("4.1.6 Transformations and Scaling")
    sk = C["skew"]
    tn = rep.next_tab(); fn = rep.next_fig()
    rep.para(f"Min-Max scaling (0-1) and standard (z-score) scaling were applied to the ten acoustic metrics and saved as "
             f"separate tables, because distance-based methods (PCA, KMeans, nearest-neighbour search) need comparable units. "
             f"Skewed fields were transformed and their skewness G1 recorded before and after (TABLE "
             f"{roman(tn)}). Duration fell from G1 = {f(sk['duration_min']['before'], 2)} to "
             f"{f(sk['duration_min']['after'], 2)} after log1p (Fig. {fn}). The genre table’s popularity score is a bounded "
             f"0-100 index with many zeros, not a count, and log1p made it more skewed (G1 {f(sk['popularity (genre table)']['before'], 2)} "
             f"to {f(sk['popularity (genre table)']['after'], 2)}), so popularity is analysed untransformed. No stream counts exist in any of the datasets.")
    rep.table(["Feature", "Transform", "G1 before", "G1 after"],
              [[k, v["transform"], f(v["before"], 3), f(v["after"], 3)] for k, v in sk.items()],
              "Skewness before and after transformation", widths_cm=[4.2, 5.5, 3, 3])
    rep.figure(FIG / "fig_02_02_duration_log_transform.png", "Duration before and after the log1p transform")
    rep.h3("4.1.7 Outlier Detection")
    io = C["iqr"]; ex = C["outlier_examples"]
    tn = rep.next_tab(); fn = rep.next_fig()
    rep.para(f"Outliers were detected in two complementary ways and flagged, never deleted, so the website can include or "
             f"exclude them. Tukey fences (Q1 - 1.5 IQR, Q3 + 1.5 IQR) were applied per feature (TABLE "
             f"{roman(tn)}), and an Isolation Forest {cite_sw('isolation')} with 200 trees and 2% "
             f"contamination was fitted on the standardised features, flagging {n(C['iso_outliers'])} tracks "
             f"({pct(C['iso_pct'])}). The flagged tracks are mainly spoken word: their mean speechiness is "
             f"{f(C['iso_outlier_mean_speechiness'], 2)} against {f(C['iso_inlier_mean_speechiness'], 2)} for the rest (Fig. {fn}).")
    rep.table(["Feature", "Lower fence", "Upper fence", "Outliers", "Share"],
              [[k, f(v["lower_fence"], 3), f(v["upper_fence"], 3), n(v["n_outliers"]), pct(v["pct"], 2)] for k, v in io.items()],
              "Tukey-fence (IQR) outliers per feature", widths_cm=[3.6, 3, 3, 3, 2.5])
    rep.figure(FIG / "fig_02_03_isolation_forest_outliers.png", "Isolation Forest outliers in the speechiness-loudness plane", 5.6)
    rep.para("Real examples confirm that the flags pick out genuine special cases rather than errors:")
    rep.bullet("Extreme tempo", f"{n(ex['tempo_over_220'])} tracks above 220 BPM, usually double-time readings of fast punk or drum-and-bass.")
    rep.bullet("Near-silent or ambient", f"{n(ex['loudness_below_-45'])} tracks quieter than -45 dB, such as field recordings, sleep sounds and silent tracks.")
    rep.bullet("Spoken word and comedy", f"{n(ex['speechiness_over_0.66'])} tracks with speechiness above 0.66 (the level Spotify associates with speech-only content), {n(ex['speechiness_over_0.9'])} of them above 0.9, mostly audiobooks, poetry and stand-up comedy.")


def proposed_42():
    rep.h2("4.2 Descriptive Statistics, Multivariate Visual Storytelling, Dimensionality Reduction and Time-Series Dynamics of Musical Characteristics", center=True)
    rep.para("This block explores the cleaned data with the statistical and visual methods of the course, moving from single features to the full multivariate space and then to time.")
    rep.h3("4.2.1 Statistical Conventions")
    tu = D["tutorial"]
    tn = rep.next_tab()
    rep.para("All summaries follow the conventions used in the course tutorials: sample variance with denominator n-1; "
             "Q1 and Q3 by Tukey’s median-of-halves (hinges); the adjusted Fisher-Pearson skewness G1; and the "
             "bias-corrected excess kurtosis G2, with Pearson kurtosis = G2 + 3. The functions live in the package module "
             "sonicatlas.stats and are unit-tested against SciPy. On the tutorial data set {4, 5, 5, 6, 6, 6, 7, 8, 9, 14} they "
             f"reproduce the expected values (TABLE {roman(tn)}). The pandas default quartiles, "
             "which interpolate linearly, are shown for comparison; the report uses the Tukey values.")
    rep.table(["Statistic", "Formula / rule", "Value"], [
        ["Mean", "sum(x) / n", f(tu["mean"], 2)], ["Median", "middle value", f(tu["median"], 2)], ["Mode", "most frequent", f(tu["mode"], 0)],
        ["Sample variance", "sum(x - mean)^2 / (n - 1)", f(tu["variance"], 3)], ["Standard deviation", "sqrt(variance)", f(tu["std"], 3)],
        ["Q1, Q3 (Tukey)", "medians of lower and upper halves", f"{f(tu['q1'], 2)}, {f(tu['q3'], 2)}"],
        ["Q1, Q3 (pandas default)", "linear interpolation", f"{f(tu['q1_pandas'], 2)}, {f(tu['q3_pandas'], 2)}"],
        ["IQR (Tukey)", "Q3 - Q1", f(tu["iqr"], 2)], ["Skewness G1", "adjusted Fisher-Pearson", f(tu["g1"], 4)],
        ["Excess kurtosis G2", "bias-corrected", f(tu["g2"], 4)], ["Pearson kurtosis", "G2 + 3", f(tu["pearson_kurtosis"], 4)],
    ], "Worked example on the tutorial data set", widths_cm=[4.5, 7, 3.5])
    rep.h3("4.2.2 Descriptive Statistics of the Audio Features")
    al = D["all"]
    tn = rep.next_tab(); fn = rep.next_fig()
    rep.para(f"TABLE {roman(tn)} gives the summary for all {n(A['a_rows'])} tracks and Fig. {fn} the "
             f"distributions. Danceability and tempo are close to symmetric. Speechiness (G1 = {f(al['speechiness']['g1'], 2)}, "
             f"G2 = {f(al['speechiness']['g2'], 1)}), liveness and duration are strongly right-skewed and heavy-tailed: most "
             f"tracks are sung studio recordings of three to five minutes, with a long tail of spoken, live and very long pieces. "
             f"Loudness is left-skewed (G1 = {f(al['loudness']['g1'], 2)}) because a minority of quiet classical and ambient "
             f"tracks pulls the tail down. Acousticness and instrumentalness are platykurtic (G2 below -1) because they pile up "
             f"near both 0 and 1.")
    feats = ["danceability", "energy", "loudness", "speechiness", "acousticness", "instrumentalness", "liveness", "valence", "tempo", "duration_min"]
    rows = [[k, f(al[k]["mean"], 3), f(al[k]["median"], 3), f(al[k]["mode"], 3), f(al[k]["variance"], 3), f(al[k]["q1"], 3),
             f(al[k]["q3"], 3), f(al[k]["iqr"], 3), f(al[k]["g1"], 2), f(al[k]["g2"], 2), f(al[k]["cv"], 2)] for k in feats]
    rep.table(["Feature", "Mean", "Median", "Mode", "Var (n-1)", "Q1", "Q3", "IQR", "G1", "G2", "CV"], rows,
              f"Descriptive statistics of audio features (n = {n(A['a_rows'])})", size=8.5,
              widths_cm=[2.9, 1.4, 1.4, 1.4, 1.5, 1.4, 1.4, 1.4, 1.2, 1.3, 1.1])
    rep.figure(FIG / "fig_03_01_feature_histograms_kde.png", "Histograms with kernel density estimates of the ten audio features", 6.4)
    shape_rows = [[k, D["shape"][k]] for k in feats]
    rep.para("The plain-English reading of each shape, generated from G1 and G2 with fixed thresholds (|G1| < 0.5 symmetric, G1 >= 1 strongly skewed; G2 > 1 heavy-tailed, G2 < -1 flat), is:")
    rep.table(["Feature", "Shape interpretation"], shape_rows, "Interpretation of distribution shapes", widths_cm=[4, 11])
    rep.h3("4.2.3 Categorical Distributions")
    fn = rep.next_fig()
    ex = D["explicit_pct_by_decade"]
    rep.para(f"Fig. {fn} shows the categorical fields. The three most common keys are {', '.join(D['top_keys'])}, which are "
             f"easy keys on guitar and piano, and {pct(D['major_pct'])} of tracks are in a major mode. The share of tracks "
             f"marked explicit stays below 1% before the 1980s, reaches {pct(ex['1990'])} in the 1990s and "
             f"{pct(ex['2010'])} in the 2010s, and is {pct(ex['2020'])} for 2020 releases, the single year available "
             f"in the 2020s.")
    rep.figure(FIG / "fig_03_02_key_mode_explicit_bars.png", "Key, mode and explicit-content distributions", 6.4)
    tn = rep.next_tab()
    dm = D["decade_means"]
    rep.para(f"Per-decade means for canonical releases (TABLE {roman(tn)}) already reveal the main "
             "historical story that Section 4.2.9 analyses formally.")
    rep.table(["Decade", "Danceability", "Energy", "Acousticness", "Valence", "Loudness (dB)", "Duration (min)"],
              [[f"{k}s", f(v["danceability"], 3), f(v["energy"], 3), f(v["acousticness"], 3), f(v["valence"], 3),
                f(v["loudness"], 2), f(v["duration_min"], 2)] for k, v in dm.items()],
              "Mean audio features by decade (canonical releases)", size=9.5)
    gm = D["genre_means"]
    tn = rep.next_tab()
    rep.para(f"TABLE {roman(tn)} gives the same summary for the eight representative genres used "
             "throughout the multivariate analysis.")
    rep.table(["Genre", "Danceability", "Energy", "Acousticness", "Valence", "Loudness (dB)", "Tempo (BPM)"],
              [[g, f(gm[g]["danceability"], 3), f(gm[g]["energy"], 3), f(gm[g]["acousticness"], 3), f(gm[g]["valence"], 3),
                f(gm[g]["loudness"], 2), f(gm[g]["tempo"], 1)] for g in sorted(gm)],
              "Mean audio features of eight representative genres", size=9.5)

    rep.h3("4.2.4 Correlation Analysis")
    pr = M["pairs"]
    fn = rep.next_fig(); tn = rep.next_tab()
    rep.para(f"Pearson and Spearman correlation matrices were computed on a 300,000-track random sample (Fig. {fn}). Three "
             f"trade-offs dominate the audio space:")
    rep.bullet("Loudness and energy", f"r = {f(pr['loudness~energy']['pearson'], 3)} (Spearman {f(pr['loudness~energy']['spearman'], 3)}); Spotify’s energy measure is built partly from loudness, so the two move together.")
    rep.bullet("Energy and acousticness", f"r = {f(pr['energy~acousticness']['pearson'], 3)}; acoustic recordings are quieter and less intense.")
    rep.bullet("Valence and danceability", f"r = {f(pr['valence~danceability']['pearson'], 3)}; happier-sounding tracks tend to be more danceable, but the link is only moderate.")
    rep.para(f"Spearman values are close to Pearson values, so the relationships are monotonic and not driven by outliers. "
             f"To connect with the tutorial formula, TABLE {roman(tn)} works through Pearson’s r for "
             f"energy and loudness on ten tracks; the hand calculation gives r = {f(M['worked_pearson']['r_hand'], 4)}, "
             f"identical to numpy ({f(M['worked_pearson']['r_numpy'], 4)}).")
    rep.figure(FIG / "fig_04_01_correlation_heatmaps.png", "Pearson and Spearman correlation heatmaps of the audio features", 6.4)
    import pandas as pd
    w = pd.read_csv(ROOT / "reports" / "tables" / "tab_04_03_worked_pearson.csv")
    rows = [[str(i + 1), f(x.energy, 3), f(x.loudness, 3), f(x.dx, 4), f(x.dy, 4), f(x["dx*dy"], 4), f(x["dx^2"], 4), f(x["dy^2"], 4)] for i, x in w.iterrows()]
    rows.append(["Sum", "", "", "", "", f(w["dx*dy"].sum(), 4), f(w["dx^2"].sum(), 4), f(w["dy^2"].sum(), 4)])
    rep.table(["Track", "Energy x", "Loudness y", "dx", "dy", "dx*dy", "dx^2", "dy^2"], rows,
              "Worked Pearson correlation by formula (energy vs loudness, 10 tracks)", size=9)

    rep.h3("4.2.5 Multivariate Visualisation of Genres")
    gf = M["genre_fingerprints"]
    fns = [rep.next_fig() + i for i in range(4)]
    rep.para(f"To see all features at once, a balanced sample of {M['balanced_per_genre']} tracks from each of the eight genres "
             f"({', '.join(FOCUS)}) was drawn so that no genre dominates the picture. In the parallel-coordinates plot "
             f"(Fig. {fns[0]}) each track is a line across seven axes, and the thick lines are genre means. Classical stands "
             f"apart with high acousticness ({f(gf['classical']['acousticness'], 2)}) and instrumentalness "
             f"({f(gf['classical']['instrumentalness'], 2)}); metal has the highest energy ({f(gf['metal']['energy'], 2)}); "
             f"hip-hop has the highest speechiness ({f(gf['hip-hop']['speechiness'], 2)}) and danceability "
             f"({f(gf['hip-hop']['danceability'], 2)}).")
    rep.figure(FIG / "fig_04_02_parallel_coordinates_genres.png", "Parallel coordinates of eight genres (balanced sample, genre means in bold)", 6.4)
    rep.para(f"Andrews curves (Fig. {fns[1]}) turn each track into a Fourier series f(t) = x1/sqrt(2) + x2 sin t + x3 cos t + ..., "
             "so tracks with similar features draw similar waves. Classical curves form a clearly separate band (deep dips near t = -2.8 and t = 1.8), while "
             "the other seven genres overlap heavily, which agrees with the parallel-coordinates view and the silhouette results in Section 4.2.7.")
    rep.figure(FIG / "fig_04_03_andrews_curves_genres.png", "Andrews curves of eight genres", 6.2)
    rep.para(f"Radar fingerprints (Fig. {fns[2]}) give each genre a compact visual signature, and violin plots (Fig. {fns[3]}) "
             f"show the full distributions behind the means. EDM and hip-hop have narrow, high danceability distributions, "
             f"while classical has the lowest energy (mean {f(gf['classical']['energy'], 2)}).")
    rep.figure(FIG / "fig_04_04_radar_genre_fingerprints.png", "Radar fingerprints of mean audio profiles per genre", 6.2)
    rep.figure(FIG / "fig_04_05_violins_by_genre.png", "Violin plots of danceability, energy and valence by genre", 6.4)

    rep.h3("4.2.6 Principal Component Analysis and SVD")
    ev = DR["evr"]; ld = DR["loadings"]
    fn = rep.next_fig(); tn = rep.next_tab()
    rep.para(f"PCA was fitted on all {n(A['a_rows'])} tracks using the ten standardised features (with log duration). "
             f"PC1 explains {pct(100 * ev[0])} of the variance and PC2 {pct(100 * ev[1])}; {DR['n90']} components are needed "
             f"to reach 90% (Fig. {fn}). The loadings in TABLE {roman(tn)} show that PC1 is an "
             f"“electric versus acoustic” axis (energy {f(ld['PC1']['energy'], 2)}, loudness {f(ld['PC1']['loudness'], 2)}, "
             f"acousticness {f(ld['PC1']['acousticness'], 2)}); PC2 contrasts long, instrumental pieces (log duration "
             f"{f(ld['PC2']['log_duration'], 2)}) with danceable, wordy tracks (danceability {f(ld['PC2']['danceability'], 2)}, "
             f"speechiness {f(ld['PC2']['speechiness'], 2)}); PC3 is dominated by liveness ({f(ld['PC3']['liveness'], 2)}). "
             f"The link to SVD was checked numerically: the squared singular values of the centred data matrix, divided by "
             f"their sum, equal PCA’s explained-variance ratios ({'match confirmed' if DR['svd_match'] else 'mismatch'}).")
    rep.figure(FIG / "fig_05_01_pca_scree_biplot.png", "PCA scree plot and biplot", 6.4)
    rep.table(["Feature", "PC1", "PC2", "PC3"], [[k, f(ld["PC1"][k], 3), f(ld["PC2"][k], 3), f(ld["PC3"][k], 3)] for k in ld["PC1"]],
              "PCA loadings of the first three components", widths_cm=[4.5, 3, 3, 3])

    rep.h3("4.2.7 Non-linear Maps: t-SNE, LLE and UMAP")
    cp = DR["comparison"]
    fn = rep.next_fig(); tn = rep.next_tab()
    rep.para(f"Non-linear methods were run on a stratified sample of {n(DR['sample_size'])} tracks (up to {DR['per_genre']} "
             f"per genre, all {A['g_genres']} genres) because their cost grows quickly with the number of points. t-SNE used "
             f"perplexity 30; LLE used 15 neighbours; UMAP was added as a bonus. Fig. {fn} compares the maps and TABLE "
             f"{roman(tn)} the metrics. t-SNE preserves local neighbourhoods best "
             f"(trustworthiness {f(cp['t-SNE']['trustworthiness_k10'], 3)}), followed by UMAP "
             f"({f(cp['UMAP']['trustworthiness_k10'], 3)}), PCA ({f(cp['PCA']['trustworthiness_k10'], 3)}) and LLE "
             f"({f(cp['LLE']['trustworthiness_k10'], 3)}). Silhouette scores computed with the eight focus genres as labels "
             f"are slightly negative for every method and even in the original space "
             f"({f(cp['Original 10-D']['silhouette_focus_genres'], 3)}): genres overlap heavily in Spotify’s feature space, "
             f"so no projection can separate them cleanly. This is an important honest finding: genre is partly a cultural "
             f"label that audio descriptors capture only in part.")
    rep.figure(FIG / "fig_05_02_tsne_lle_pca_maps.png", "PCA, t-SNE, LLE and UMAP maps of a stratified 10,000-track sample", 6.5)
    rep.table(["Method", "Runtime (s)", "Trustworthiness (k=10)", "Silhouette (8 genres)"],
              [[k, f(v["runtime_s"], 2), f(v["trustworthiness_k10"], 3), f(v["silhouette_focus_genres"], 3)] for k, v in cp.items()],
              "Comparison of dimensionality-reduction methods", widths_cm=[3.6, 3, 4.2, 4.2])
    fn = rep.next_fig()
    rep.para(f"Colouring a decade-stratified t-SNE map (1,000 canonical tracks per decade, 1950s-2010s) by release decade "
             f"gives a silhouette of {f(DR['decade_tsne_silhouette'], 3)} (Fig. {fn}): eras blend into each other gradually "
             "rather than forming islands, consistent with the slow trends in Section 4.2.9.")
    rep.figure(FIG / "fig_05_03_tsne_by_decade.png", "t-SNE map coloured by release decade", 5.0)

    rep.h3("4.2.8 Acoustic Clusters")
    cl = DR["clusters"]
    fn = rep.next_fig(); tn = rep.next_tab()
    rep.para(f"KMeans was run on the first {DR['n90']} principal components for k = 2 to 10 (Fig. {fn}). The silhouette is "
             f"highest at k = {DR['best_k']} ({f(DR['silhouette_by_k'][0], 3)}), which only splits acoustic from electric "
             f"music; k = {DR['k_used']} sits on the next plateau and produces groups that can be described in plain words, so "
             f"it was chosen. Each cluster was named from its centroid and the most frequent genres of genre-table tracks "
             f"assigned to it (TABLE {roman(tn)}). HDBSCAN on a 20,000-track subsample found only "
             f"{DR['hdbscan_clusters']} dense clusters and labelled {pct(DR['hdbscan_noise_pct'])} of points as noise, "
             "confirming that the audio space is a continuum with a few dense cores rather than well-separated islands.")
    rep.figure(FIG / "fig_05_04_kmeans_elbow_silhouette.png", "Elbow and silhouette curves for choosing k", 5.8)
    rep.table(["Cluster name", "Share", "Top genres", "Energy", "Acoust.", "Speech.", "Instr.", "Dance."],
              [[v["name"], pct(v["share_pct"]), v["top_genres"], f(v["energy"], 2), f(v["acousticness"], 2), f(v["speechiness"], 2),
                f(v["instrumentalness"], 2), f(v["danceability"], 2)] for v in cl.values()],
              f"KMeans (k = {DR['k_used']}) acoustic clusters with plain-language names", size=8.5,
              widths_cm=[3.6, 1.4, 4.0, 1.4, 1.5, 1.5, 1.4, 1.5])
    rep.figure(FIG / "fig_05_05_clusters_on_tsne.png", "Named KMeans clusters on the t-SNE map", 5.2)

    rep.h3("4.2.9 Time-Series Dynamics and Decade Trends")
    fl = TS["first_last"]; cps = TS["change_points"]; st = TS["stl"]; ar = TS["arima"]
    fns = [rep.next_fig() + i for i in range(6)]
    rep.para(f"Years with at least 500 canonical releases ({TS['year_window'][0]}-{TS['year_window'][1]}, {n(TS['tracks_in_window'])} "
             f"tracks) form the analysis window (Fig. {fns[0]}); the catalogue is heavily weighted towards recent years, so "
             "early decades rest on smaller samples.")
    rep.figure(FIG / "fig_06_01_tracks_per_year.png", "Number of canonical tracks per release year", 6.2)
    rep.para(f"Yearly means with 5-year rolling averages (Fig. {fns[1]}) show clear long-run movements. From "
             f"{TS['year_window'][0]} to {TS['year_window'][1]} mean loudness rose from {f(fl['loudness'][0], 1)} dB to "
             f"{f(fl['loudness'][1], 1)} dB, energy from {f(fl['energy'][0], 2)} to {f(fl['energy'][1], 2)} and danceability "
             f"from {f(fl['danceability'][0], 2)} to {f(fl['danceability'][1], 2)}, while acousticness fell from "
             f"{f(fl['acousticness'][0], 2)} to {f(fl['acousticness'][1], 2)} and average duration shrank from "
             f"{f(fl['duration_min'][0], 2)} to {f(fl['duration_min'][1], 2)} minutes. Valence moved much less "
             f"({f(fl['valence'][0], 2)} to {f(fl['valence'][1], 2)}).")
    rep.figure(FIG / "fig_06_02_yearly_trends_rolling.png", "Yearly means and 5-year rolling means of six audio features", 6.5)
    rep.para(f"Change points were located with the PELT algorithm and an RBF cost from the ruptures library {cite_sw('ruptures')} "
             f"(Fig. {fns[2]}). For loudness the breaks fall at {', '.join(map(str, cps['loudness']))}; the 1989 and 1999 breaks "
             f"match the period that audio engineers call the “loudness war”, when mastering for CDs pushed average "
             f"levels up. Energy and acousticness both break in {cps['energy'][-1]}. The break at {cps['loudness'][0]} appears "
             "in every feature and mainly separates the thin late-1950s sample from the rest, so it should be read with care.")
    rep.figure(FIG / "fig_06_03_change_points.png", "Change points in yearly loudness and acousticness", 6.2)
    rep.para(f"Seasonality was examined with STL decomposition {cite_sw('stl')} of monthly means (1990-2020) using only "
             f"day-precision dates. Releases dated 1 January were excluded because {pct(TS['jan1_share_pct'])} of day-precision "
             f"dates fall on that day, a typical placeholder. The seasonal component of valence (Fig. {fns[3]}) has strength "
             f"{f(st['valence']['seasonal_strength'], 2)} on a 0-1 scale, peaking in month {st['valence']['peak_month']} and "
             f"lowest in month {st['valence']['low_month']}; energy behaves similarly (strength {f(st['energy']['seasonal_strength'], 2)}). "
             "These are modest effects: summer releases sound slightly brighter, but the trend dominates.")
    rep.figure(FIG / "fig_06_04_stl_valence_monthly.png", "STL decomposition of monthly mean valence (1990-2020)", 6.0)
    rep.para(f"A simple ARIMA(1,1,1) model was trained on yearly loudness up to {ar['train_end']} and tested on the next "
             f"{ar['test_years']} years (Fig. {fns[4]}). Its mean absolute error is {f(ar['mae'], 2)} dB, almost the same as "
             f"a naive forecast that repeats the last training value ({f(ar['mae_naive'], 2)} dB). Loudness levelled off after "
             "about 2010, so a trend model gains little: an honest reminder that past trends do not guarantee future ones.")
    rep.figure(FIG / "fig_06_05_arima_loudness_forecast.png", "ARIMA forecast of yearly loudness against held-out years", 6.0)
    dm = TS["decade_means"]
    rep.para(f"Fig. {fns[5]} condenses the decade story. The 1950s were acoustic and calm (acousticness "
             f"{f(dm['1950']['acousticness'], 2)}, energy {f(dm['1950']['energy'], 2)}). The 1960s and 1970s brought the "
             f"electric era: by the 1970s energy had reached {f(dm['1970']['energy'], 2)} and acousticness had dropped to "
             f"{f(dm['1970']['acousticness'], 2)}. The 1980s held steady, and the 1990s decade mean actually dips (energy "
             f"{f(dm['1990']['energy'], 2)}, acousticness {f(dm['1990']['acousticness'], 2)}), which reflects the mix of "
             f"releases catalogued for that decade rather than a change in mainstream production. From the 2000s the rise "
             f"resumes: the 2010s average energy {f(dm['2010']['energy'], 2)} and loudness {f(dm['2010']['loudness'], 1)} dB, "
             f"and the 2020 releases are the most danceable ({f(dm['2020']['danceability'], 2)}) and least acoustic "
             f"({f(dm['2020']['acousticness'], 2)}) in the data.")
    rep.figure(FIG / "fig_06_06_decade_story.png", "Decade means of danceability, energy, acousticness and valence", 6.0)


def proposed_43():
    rep.h2("4.3 The SonicAtlas Interactive Query Platform: Architecture, Text Mining, Hypothesis Testing and Evaluation Framework", center=True)
    rep.para("This block describes the platform that delivers the analysis to users, the text and statistical tests that "
             "complete the exploration, and the framework used to evaluate the whole system.")
    rep.h3("4.3.1 System Architecture")
    fn = rep.next_fig()
    rep.para(f"SonicAtlas separates heavy computation from interaction (Fig. {fn}). All cleaning, embeddings, clustering and "
             f"statistics run offline in Python and are exported as compact Parquet files and a single results.json. The "
             f"website is a static bundle: when it opens, DuckDB-WASM {cite_sw('duckdb-wasm')} starts in a Web Worker, fetches "
             f"the Parquet file over HTTP and answers SQL queries locally, so no data leaves the user’s machine and the "
             f"site can be hosted for free on GitHub Pages.")
    rep.figure(FIG / "fig_00_01_system_architecture.png", "System architecture of SonicAtlas", 6.4)
    rep.h3("4.3.2 Platform Modules")
    tn = rep.next_tab()
    rep.para(f"TABLE {roman(tn)} lists the modules of the platform and their purpose. Charts follow the grammar-of-graphics "
             f"approach of Vega-Lite {cite('satyanarayan2017vegalite')}, and the Collaboration Network module is computed with "
             f"NetworkX {cite('hagberg2008networkx')} and Louvain community detection {cite('blondel2008louvain')}.")
    status = (WEB or {}).get("page_status", {})
    mods = [("Home", "headline statistics, mini acoustic map, three key findings"),
            ("Query Studio", "visual filter builder, SQL console with examples and history, plain-English query box that shows its SQL, export to CSV/JSON"),
            ("Acoustic Map", "PCA / t-SNE / LLE / UMAP switch, colour by genre, cluster or feature, hover details, click for similar tracks"),
            ("Sonic Twins", "nearest tracks by cosine similarity on scaled features with a radar comparison"),
            ("Compare", "two or three groups side by side with G1/G2/IQR table, Welch t / Mann-Whitney / ANOVA and effect size"),
            ("Time Machine", "yearly trends, STL components and change points per feature"),
            ("Genre Fingerprints", "parallel coordinates, radar profiles and Andrews curves"),
            ("Lyrics Lab", "n-grams, sentiment versus valence, vocabulary diversity"),
            ("Collaboration Network", "artist graph with communities"),
            ("Data Profile", "embedded automated profiling reports and quick profile of any query"),
            ("Methodology", "sources, cleaning steps, statistical conventions, limitations")]
    hdr = ["Module", "Purpose"] + (["Status"] if status else [])
    rows = [[m, p] + ([status.get(m, "planned")] if status else []) for m, p in mods]
    rep.table(hdr, rows, "Modules of the SonicAtlas platform", widths_cm=[3.4, 10.0] + ([2.6] if status else []))
    if not WEB:
        rep.para("At the time of this report version the analysis pipeline and data exports are complete and the web "
                 "front-end is being implemented; Section 6 gives the repository where the working build is published.")
    else:
        rep.para(WEB.get("report_paragraph", ""))

    rep.h3("4.3.3 Hypothesis Testing")
    an = T["anova"]; kw = T["kruskal"]; cg = T["chi2_mode_genre"]; cd = T["chi2_mode_decade"]
    tn = rep.next_tab()
    rep.para("Three formal tests were carried out in the step-by-step style of the course tutorials, at significance level alpha = 0.05.")
    rep.para("**Test 1 - One-way ANOVA on danceability across eight genres.** H0: all eight genre means are equal; H1: at "
             f"least one differs. F = MS_between / MS_within = [SS_B / (k - 1)] / [SS_W / (N - k)]. With k = {an['k']} and "
             f"N = {n(an['N'])}, SS_B = {f(an['SSB'], 2)} and SS_W = {f(an['SSW'], 2)}, giving F({an['df_between']}, "
             f"{n(an['df_within'])}) = {f(an['F'], 1)} and {pval(an['p'])}. Decision: reject H0. Effect size eta squared = "
             f"{f(an['eta_squared'], 3)}, a large effect: genre explains about {pct(100 * an['eta_squared'], 0)} of the "
             f"variation in danceability. Tukey’s HSD post-hoc test found {T['tukey_sig_pairs']} of {T['tukey_pairs']} genre "
             f"pairs significantly different (TABLE {roman(tn)} lists the group means).")
    rep.table(["Genre", "n", "Mean danceability"], [[g, n(an["group_n"][g]), f(an["group_means"][g], 3)] for g in FOCUS],
              "Group sizes and means used in the ANOVA", widths_cm=[4, 3, 4])
    rep.para(f"**Test 2 - Kruskal-Wallis check.** Danceability is bounded and Levene’s test rejects equal variances "
             f"(W = {f(T['levene']['W'], 2)}, {pval(T['levene']['p'])}), so the rank-based Kruskal-Wallis test was run as a "
             f"robustness check: H = {f(kw['H'], 1)}, {pval(kw['p'])}, epsilon squared = {f(kw['epsilon_squared'], 3)}. The "
             "conclusion of the ANOVA is unchanged.")
    tn = rep.next_tab()
    rep.para(f"**Test 3 - Chi-square test of independence between mode and genre.** H0: mode (major/minor) is independent "
             f"of genre. chi2 = sum (O - E)^2 / E with E = row total x column total / N. Result: chi2({cg['dof']}) = "
             f"{f(cg['chi2'], 1)}, {pval(cg['p'])}, Cramer’s V = {f(cg['cramers_v'], 3)} (moderate). Country is the most "
             f"major-key genre ({pct(cg['major_pct']['country'])}) and hip-hop the least ({pct(cg['major_pct']['hip-hop'])}) "
             f"(TABLE {roman(tn)}). Mode versus decade on the full catalogue is also significant "
             f"(chi2({cd['dof']}) = {f(cd['chi2'], 0)}, {pval(cd['p'])}) but with a tiny effect (V = {f(cd['cramers_v'], 3)}): "
             f"the major share slipped from {pct(cd['major_pct']['1960'])} in the 1960s to {pct(cd['major_pct']['2010'])} in "
             "the 2010s. With over a million rows almost any difference is significant, so effect sizes are the better guide.")
    rep.table(["Genre", "Major-key share"], [[g, pct(cg["major_pct"][g])] for g in FOCUS], "Share of major-key tracks per genre", widths_cm=[4, 4])

    rep.h3("4.3.4 Feature Selection: Filter versus Wrapper Methods")
    fr = T["feature_ranks"]
    fn = rep.next_fig(); tn = rep.next_tab()
    rep.para(f"To find which audio features best separate the eight genres, four filter criteria (variance after Min-Max "
             f"scaling, ANOVA F, mutual information, and pairwise correlation as a redundancy check) were compared with a "
             f"wrapper method, recursive feature elimination (RFE) around multinomial logistic regression (TABLE "
             f"{roman(tn)}, Fig. {fn}). Acousticness ranks first or second in every method, and "
             f"loudness ranks in the top two for ANOVA F, mutual information and RFE. Top three by ANOVA F: "
             f"{', '.join(T['top3_anova'])}; by mutual information: {', '.join(T['top3_mi'])}; by RFE: "
             f"{', '.join(T['top3_rfe'])}. Liveness is last in three methods, so it carries little genre information.")
    feats = list(fr["anova_F"].keys())
    rep.table(["Feature", "Variance rank", "ANOVA F rank", "Mutual info rank", "RFE rank"],
              [[k, fr["variance"][k], fr["anova_F"][k], fr["mutual_info"][k], fr["rfe"][k]] for k in feats],
              "Feature rankings from filter and wrapper methods (1 = most useful)", widths_cm=[3.8, 2.8, 2.8, 3, 2.6])
    rep.figure(FIG / "fig_08_01_feature_selection_ranks.png", "Feature importance ranks from filter and wrapper methods", 6.2)

    rep.h3("4.3.5 Lyrics Text Mining")
    if LY:
        for ptxt in LY.get("report_paragraphs", []):
            rep.para(ptxt)
        k = rep.next_fig()
        rep.para(f"Sentiment scores come from the VADER lexicon {cite('hutto2014vader')} and the TextBlob library {cite('textblob')}. Fig. {k} plots lyric sentiment against audio valence and compares the two sentiment tools; Fig. {k + 1} "
                 "breaks sentiment down by genre tag and decade and shows vocabulary diversity per tag.")
        for fg, cap in LY.get("report_figures", []):
            rep.figure(FIG / fg, cap, 6.0)
    else:
        rep.para("The lyrics module samples English songs from the Genius dataset, cleans the text, and computes n-gram "
                 "frequencies, type-token ratio, repetition and sentiment with VADER and TextBlob, then compares lyric "
                 "sentiment with audio valence for the songs that match the audio catalogue. Results are added in the next "
                 "version of this report.")
    rep.h3("4.3.6 Artist Collaboration Network")
    if NET:
        fn = rep.next_fig()
        td = ", ".join(f"{a} ({d})" for a, d in NET["top_degree"][:5])
        tb = ", ".join(a for a, _ in NET["top_betweenness"][:3])
        rep.para(f"Artists who share a canonical track were linked to form a collaboration graph with {n(NET['nodes'])} artists "
                 f"and {n(NET['edges'])} edges, built from {n(NET['collab_tracks'])} multi-artist tracks. The graph is very "
                 f"sparse (density {NET['density']:.1e}) but {pct(NET['giant_share_pct'])} of artists sit in one giant connected "
                 f"component. Louvain community detection {cite('blondel2008louvain')} found {NET['communities']} communities "
                 f"with a high modularity of {f(NET['modularity'], 2)}, meaning collaborations stay mostly inside scenes. The "
                 f"highest-degree nodes are {td}: classical recordings credit the composer together with many performers, and "
                 f"placeholder names such as “Traditional” and “Anonymous” act as hubs. Betweenness (approximated "
                 f"with {NET['betweenness_samples']} sampled sources) is highest for {tb}, which bridge otherwise separate "
                 f"communities. Fig. {fn} shows the largest connected group among the 300 most connected artists; the same subgraph is explorable on the website.")
        rep.figure(FIG / "fig_10_01_collaboration_network.png", "Largest connected component of the 300 most connected artists, coloured by Louvain community", 5.8)
    rep.h3("4.3.7 Automated EDA Tools")
    if PR:
        for ptxt in PR.get("report_paragraphs", []):
            rep.para(ptxt)
        if PR.get("table"):
            rep.para(f"TABLE {roman(rep.next_tab())} summarises the comparison.")
            rep.table(PR["table"]["header"], PR["table"]["rows"], PR["table"]["caption"], size=9, widths_cm=[2.8, 1.9, 2.6, 4.3, 4.3])
    else:
        rep.para(f"YData-Profiling, Sweetviz and D-Tale, the three tools in the course handout, are run on the same sample "
                 f"and compared with the manual EDA on what they detect, what they miss, runtime and usability.")
    rep.h3("4.3.8 Evaluation Framework")
    tn = rep.next_tab()
    rep.para(f"Each part of the system is evaluated with an explicit, reproducible criterion (TABLE {roman(tn)}).")
    rows = [["Data quality", "MICE RMSE vs mean imputation on 5% masked values", f"MICE better for all 3 columns (energy {pct(C['mice_validation']['energy']['improvement_pct'])})"],
            ["Statistics code", "unit tests against SciPy and the tutorial data", (r.get("qa", {}).get("pytest") or "7 of 7 tests pass")],
            ["Embeddings", "trustworthiness (k = 10), silhouette, runtime", f"t-SNE {f(DR['comparison']['t-SNE']['trustworthiness_k10'], 3)} best"],
            ["Clustering", "silhouette and elbow over k = 2-10", f"k = {DR['k_used']} chosen, named clusters"],
            ["Forecasting", "hold-out MAE vs naive baseline", f"{f(TS['arima']['mae'], 2)} vs {f(TS['arima']['mae_naive'], 2)} dB"],
            ["Tests", "p-value plus effect size", f"eta squared {f(T['anova']['eta_squared'], 3)}, V {f(T['chi2_mode_genre']['cramers_v'], 3)}"],
            ["Website", "query latency on lite data, Lighthouse, end-to-end tests", (WEB or {}).get("eval_summary", "measured in the final version")]]
    rep.table(["Component", "Criterion", "Result"], rows, "Evaluation framework and results", widths_cm=[3.2, 6.4, 6.4])
    rep.h3("4.3.9 Summary and Method Selection Guide")
    tn = rep.next_tab()
    rep.para(f"TABLE {roman(tn)} summarises which method answered which question and when a reader should choose it.")
    guide = [["Describe one feature", "Mean/median, Tukey IQR, G1, G2", "Always first; use median and IQR for skewed fields such as speechiness"],
             ["Repair impossible values", "MICE (IterativeImputer)", "When other columns predict the gap; validate by masking"],
             ["Find unusual tracks", "Tukey fences + Isolation Forest", "Fences per feature; Isolation Forest for multivariate oddities"],
             ["Relate two features", "Pearson / Spearman", "Spearman when relationships are monotonic but not linear"],
             ["Compare genres visually", "Parallel coordinates, radar, Andrews", "Use a balanced sample to avoid clutter"],
             ["Global structure", "PCA / SVD", "Fast, interpretable loadings, works on all rows"],
             ["Local neighbourhoods", "t-SNE / UMAP", "Best trustworthiness; distances between far clusters are not meaningful"],
             ["Manifold unfolding", "LLE", "Only for smooth manifolds; weakest here"],
             ["Group tracks", "KMeans / HDBSCAN", "KMeans for nameable groups; HDBSCAN shows density and noise"],
             ["Trends over time", "Rolling means, STL, PELT, ARIMA", "Use canonical releases; check against a naive forecast"],
             ["Test differences", "ANOVA + Tukey, Kruskal-Wallis, chi-square", "Report effect sizes; large n makes tiny effects significant"],
             ["Rank features", "ANOVA F, MI, RFE", "Agreement across methods signals robust features"]]
    rep.table(["Question", "Method", "When to use it"], guide, "Method selection guide", widths_cm=[3.8, 4.6, 7.6], size=9.5)


def conclusion():
    rep.h1("5. Conclusion")
    fl = TS["first_last"]
    rep.para(f"This project set out to understand how the sound of recorded music differs across genres and how it has "
             f"changed across eras, and to make that understanding explorable by anyone. A reproducible pipeline cleaned "
             f"{n(A['a_rows'])} Spotify tracks, repaired {n(C['tempo_zero_imputed'])} impossible tempos with validated MICE "
             f"imputation, marked {n(C['reissue_rows'])} re-issues so they do not distort trends, and flagged "
             f"{n(C['iso_outliers'])} multivariate outliers instead of deleting them.")
    rep.para(f"The exploration produced clear, quantified findings. Energy is almost a restatement of loudness "
             f"(r = {f(M['pairs']['loudness~energy']['pearson'], 2)}) and the opposite of acousticness "
             f"(r = {f(M['pairs']['energy~acousticness']['pearson'], 2)}). Over {TS['year_window'][0]}-{TS['year_window'][1]} "
             f"music became about {f(fl['loudness'][1] - fl['loudness'][0], 1)} dB louder, more energetic "
             f"({f(fl['energy'][0], 2)} to {f(fl['energy'][1], 2)}) and far less acoustic ({f(fl['acousticness'][0], 2)} to "
             f"{f(fl['acousticness'][1], 2)}), with loudness change points in {', '.join(map(str, TS['change_points']['loudness'][1:]))}. "
             f"Genre explains {pct(100 * T['anova']['eta_squared'], 0)} of the variation in danceability, and acousticness and "
             f"loudness are the most useful features for telling genres apart.")
    rep.para(f"The dimensionality-reduction comparison showed that t-SNE keeps local neighbourhoods best "
             f"(trustworthiness {f(DR['comparison']['t-SNE']['trustworthiness_k10'], 2)}) while PCA gives the most "
             f"interpretable global axes, and that genres overlap in every projection. {DR['k_used']} named acoustic clusters "
             f"({', '.join(v['name'] for v in DR['clusters'].values())}) give users a vocabulary for "
             f"navigating the map. SonicAtlas packages these results in a serverless website where SQL runs inside the browser.")
    rep.para(f"The work has limitations. Spotify’s descriptors are outputs of proprietary models, not direct "
             f"measurements; the catalogue over-represents recent releases; the genre and audio tables matched for only "
             f"{pct(C['join']['match_rate_pct'])} of genre-labelled tracks, so genre-by-decade questions rest on "
             f"{n(C['genre_tracks_with_year'])} tracks; the lyrics analysis is limited to English songs; and the 2020s "
             "contain a single year.")
    rep.para("Future work includes adding newer open audio descriptors (for example from AcousticBrainz-style feature "
             "extractors) to replace the closed Spotify features, multilingual lyrics analysis, learned audio embeddings, and "
             "user studies that measure how quickly students and researchers answer questions with SonicAtlas compared with "
             "notebooks.")


def outcome():
    rep.h1("6. Project Outcome")
    rep.para("**(a) Research paper.** A six-page paper in the official IEEE conference format, titled “" + TITLE +
             "”, has been written from the results of this project. Status: prepared for submission. [Paper Link]")
    rep.para("**(b) Target conferences.** Three suitable venues:")
    for v in lit.get("venues", [])[:3]:
        p = rep.bullet(v["name"], f"{v.get('description', '')} Dates: {v.get('dates_or_deadline', 'see website')}. Website: ")
        add_hyperlink(p, v["url"], v["url"])
    rep.para("**(c) Software and data.**")
    rep.bullet("GitHub repository", L["github"])
    rep.bullet("Live website", L["live"])
    rep.bullet("Demo video", L["video"])
    rep.bullet("Key features", "in-browser SQL (DuckDB-WASM), filters and plain-English queries, switchable embedding maps, "
               "group comparison with live tests, time-series explorer, shareable URLs.")


def references():
    rep.h1("7. References")
    for i, p in enumerate(REFS, 1):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.left_indent = Cm(1.0); para.paragraph_format.first_line_indent = Cm(-1.0)
        para.paragraph_format.line_spacing = 1.15; para.paragraph_format.space_after = Pt(6)
        run = para.add_run(f"[{i}]\t{ref_text(p)} "); run.font.size = Pt(11)
        if p.get("url"):
            run2 = para.add_run("Available: "); run2.font.size = Pt(11)
            add_hyperlink(para, p["url"], p["url"], size=10)


def build(pages=None) -> Path:
    global rep, doc
    rep = Report(TITLE, "Sharvesh Prabhu (25BDS0318); Nithin S (25BDS0308)", f"{COURSE_CODE} {COURSE} - Project Report, Team 05")
    doc = rep.doc
    cover(); toc(pages); abstract(); introduction(); literature(); gaps_objectives()
    proposed_41(); proposed_42(); proposed_43(); conclusion(); outcome(); references()
    out = OUT / f"{PREFIX}_Project_Report.docx"
    OUT.mkdir(exist_ok=True)
    rep.save(out)
    return out


def find_pages(pdf: Path) -> dict[str, int]:
    """Map each TOC entry to the first page (after the TOC) whose text contains the heading."""
    d = fitz.open(pdf)
    texts = [re.sub(r"\s+", " ", pg.get_text()) for pg in d]
    pages = {}
    for num, text, main in TOC_ITEMS:
        key = text if main else text.split(" ", 1)[0]
        needle = (num + " " + text).upper() if main else text[:40]
        for i in range(2, len(texts)):
            hay = texts[i].upper() if main else texts[i]
            if needle in hay:
                pages[key] = i + 1
                break
    return pages


if __name__ == "__main__":
    docx = build()
    pdf = docx.with_suffix(".pdf")
    to_pdf_word(docx, pdf)
    pages = find_pages(pdf)
    for _ in range(2):
        docx = build(pages)
        to_pdf_word(docx, pdf)
        new = find_pages(pdf)
        if new == pages:
            break
        pages = new
    missing = [t for _, t, m in TOC_ITEMS if (t if m else t.split(" ", 1)[0]) not in pages]
    print("pages:", fitz.open(pdf).page_count, "toc:", pages, "missing:", missing)
    json.dump({"pages": fitz.open(pdf).page_count, "toc": pages}, open(ROOT / "reports" / "report_pages.json", "w"), indent=1)
