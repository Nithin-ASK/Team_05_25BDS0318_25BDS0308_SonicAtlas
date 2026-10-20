"""MANUAL_STEPS_FOR_ME (pdf + txt) and the one-page Viva Questions sheet."""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from common import OUT, PREFIX, R, f, links, n, pct  # noqa: E402
from docx_helpers import Report, to_pdf_word  # noqa: E402

r = R(); L = links()
A, C, T, DR, TS, M = r["audit"], r["cleaning"], r["tests"], r["dimred"], r["timeseries"], r["multivariate"]
fl = TS["first_last"]
PF = PREFIX

SECTIONS = [
 ("A. Check the cover page (today, 5 minutes)", [
   "Open deliverables/Project_Report.docx in Word. Look at page 1.",
   "School name: the report says SCHOOL OF COMPUTER SCIENCE AND ENGINEERING. If your school is different, change it, then File > Save.",
   "Faculty name: search for ____ (Ctrl+F) in the report and in the PPT title slides. Type the faculty name.",
   "VIT logo: download the logo from the official VIT website or your VIT login page and save it as a PNG. Click the box on the cover, delete the text '[Paste official VIT logo here]', then Insert > Pictures > This Device and choose the PNG. Make it about 4 cm wide.",
   "Then File > Export > Create PDF/XPS to make a new PDF. Check that the Table of Contents page numbers still match (the TOC is a table, not automatic: if pages moved, fix the numbers by hand).",
   "Tell Claude to re-run 'make report' instead if you want the logo added by script: give the PNG path and the faculty name.",
 ]),
 ("B. Send the three emails (when: Review-2 today; Review-3 and final when the faculty asks)", [
   "Open deliverables/Email_Drafts.txt. Each email has a subject line, a body and a list of attachment names.",
   "In Gmail or Outlook click Compose. Copy the To address (faculty email), the subject and the body from the draft. Write the faculty email in To. Add Sharvesh and Nithin in Cc.",
   "Click the paperclip and attach exactly the files listed. Review-2 email: Review2_PPT.pptx and its PDF. Check the total size (limit about 25 MB). If too big, attach only the PDFs and paste the GitHub link.",
   "The links are already filled in. Read the email once, then press Send. Nobody else can press this button for you.",
   "Keep the Sent mail as proof of submission.",
 ]),
 ("C. Submit the case study report (Lab Assessment-7, last date 20.10.2026)", [
   "Open the VIT student portal (the same place where you did the other lab assessments) and log in with your own password.",
   "Go to the course BACSE301 > Lab Assessment-7 > Upload.",
   f"Upload {PF}_Project_Report.pdf. Upload the .docx only if the portal asks for it. Keep the file name exactly as it is: it starts with {PF}.",
   "Press Submit. On the success page take a screenshot and save it in a folder called 'proofs'. This is your receipt.",
 ]),
 ("D. Plagiarism: what the file in deliverables really is", [
   f"{PF}_Plagiarism_Report.pdf is OUR OWN self-check. It is NOT a Turnitin report and must never be shown as one.",
   "Our self-check compared the report and paper with 57 sources and searched 90 sentences on the web. It found about 0.18% overlap in the report and 0% in the paper. An official tool may give different numbers.",
   "For the official check, ask the faculty or the VIT library: 'Please run a Turnitin similarity check on our report.' Send them the PDF. Save the Turnitin report they send you.",
   "When you get it, put it in the folder 'proofs' and add its first page in the Outcome Proof, in the box 'Official Turnitin report from VIT' (section Pending Proofs).",
   "WARNING: never paste your report or paper into free online plagiarism checkers. Many of them keep the text. Later the official Turnitin may then say your own text is copied.",
 ]),
 ("E. The IEEE paper and the project outcome", [
   "The paper is 'prepared for submission'. It is NOT submitted, accepted or published. Never write 'published' or 'accepted' until you have the email that proves it.",
   "Choose a conference. We found IEEE CONECCT 2027 (Bangalore, 29-31 July 2027). Its paper deadline is listed as 15 December 2026 (checked on 20 Oct 2026). Dates change: open https://www.ieee-conecct.org.ieeebangalore.org/ and check the deadline and the paper rules again before you decide.",
   "Submit only on the conference's own website (usually EDAS or Microsoft CMT). Create the account yourself, upload deliverables/IEEE_Paper.pdf, and add both authors with their VIT emails.",
   "IEEE PDF eXpress: go to ieee-pdf-express.org, create an account, enter the Conference ID that the conference gives you, and upload the PDF. If it says 'fail', send us the message and we fix the file.",
   "Keep the submission receipt email and the paper ID. When you receive them, add them to the Outcome Proof in the empty boxes 'Conference submission acknowledgement' and 'Acceptance e-mail'. Do not make these documents yourself.",
 ]),
 ("F. Team Sharvesh and Nithin", [
   "Nithin: send Sharvesh the GitHub link https://github.com/Nithin-ASK/Team_05_25BDS0318_25BDS0308_SonicAtlas and the live link https://nithin-ask.github.io/Team_05_25BDS0318_25BDS0308_SonicAtlas/ so both can use and run the website.",
   "Both of you: read the three 2-minute scripts below and the Viva Questions sheet. Practise once aloud.",
   "Sharvesh's email address is still '<email>' in the paper and documents. Give it to Claude to fill, or type it by hand.",
 ]),
]

SCRIPTS = [
 ("Review-2 script (2 minutes)",
  "We are Team 05. SonicAtlas asks how the sound of music changes across genres and eras. We cleaned about 1.2 million public Spotify tracks. "
  "Some tracks had a tempo of zero, which is impossible, so we repaired them with MICE. To test MICE we hid 5 percent of known values, and its error was up to 57 percent lower than simply filling the average. "
  "We flagged outliers and re-releases instead of deleting them. Our statistics follow our course: n minus 1 variance, Tukey quartiles, G1 and G2. "
  "Energy rises with loudness and falls with acousticness. PCA, t-SNE and LLE show that genres overlap, and t-SNE keeps local structure best. "
  "Next we finish the website, the lyrics and network analysis, and the paper."),
 ("Review-3 script (2 minutes)",
  f"Since Review-2 we finished everything. Music became about {f(fl['loudness'][1] - fl['loudness'][0], 1)} decibels louder, with change points in 1989 and 1999, the CD loudness-war years. "
  f"Genre explains {pct(100 * T['anova']['eta_squared'], 0)} of the variation in danceability. Lyric mood barely follows musical mood. "
  "The website runs SQL inside the browser with DuckDB, so the data stays private and hosting is free. It passes 65 unit tests and 8 end-to-end tests, answers a query in a few milliseconds, and scores 98 to 100 in Lighthouse. "
  "The IEEE paper is prepared for submission."),
 ("Final script (2 minutes)",
  "Open music data is messy and usually shown in static charts. We built a documented pipeline over 1.2 million tracks, 114 genres and 161 thousand English lyrics. "
  "We used the course methods: descriptive statistics, correlation, PCA, SVD, t-SNE, LLE, STL, change points, ANOVA and chi-square. "
  "We delivered a serverless website where anyone can query, map and compare the data. Our five findings: music got louder and less acoustic; energy follows loudness; genre explains 42 percent of danceability; t-SNE is the best map; and lyric sentiment is almost independent of audio valence. "
  f"Our limits: Spotify features are model-made, the catalogue favours recent years, only {pct(C['join']['match_rate_pct'])} of the genre data joins to the audio table, and lyrics are English only."),
]

TICK = [
 "File names start with Team_05_25BDS0318_25BDS0308.",
 "Faculty name and school name are correct on the cover and title slides.",
 "VIT logo is on the report cover.",
 "No [..Link..] placeholder is left (links are filled in); the words <email> and ____ are replaced by real values (search with Ctrl+F).",
 "Every PPT has at most 10 slides.",
 "The report PDF opens and the Table of Contents page numbers are right.",
 "The attachments in the email match the file names in the draft.",
 "The live website opens in a new browser window and Query Studio runs a query.",
 "You did not write 'published' or 'accepted' anywhere without proof.",
 "You saved the screenshot or email proving each submission.",
]

VIVA = [
 ("What is the project in one line?", "SonicAtlas explores how Spotify audio features and lyrics differ across genres and eras, and gives a website to query them in the browser."),
 ("Where does the data come from?", "Three public Kaggle datasets: Spotify 1.2M+ songs, Spotify tracks with 114 genres, and Genius lyrics. The live Spotify API is closed for new apps."),
 ("What is MICE and why did you use it?", "Multiple Imputation by Chained Equations: each column with gaps is predicted from the others, in turns. We used it for the impossible 0 BPM tempos."),
 ("How did you know MICE worked?", "We hid 5% of known values, imputed them, and compared the error with mean filling. MICE was better for energy, valence and tempo."),
 ("Why log transform duration?", "Duration is strongly right-skewed (G1 about 8). log1p brings G1 to about 0.2, so distance methods are not dominated by very long tracks."),
 ("What are G1 and G2?", "G1 is the adjusted Fisher-Pearson skewness; G2 is the bias-corrected excess kurtosis. Pearson kurtosis is G2 plus 3."),
 ("Why Tukey quartiles?", "It is the method used in our tutorials: split the data at the median and take the median of each half. We also show the pandas value; on a million rows they agree."),
 ("Why not delete outliers?", "Many are real: spoken word, ambient, very fast songs. We flag them (Tukey fences and Isolation Forest) so users can include or exclude them."),
 ("Pearson vs Spearman?", "Pearson measures linear relation, Spearman monotonic relation using ranks. Both agree here, so outliers are not driving the results."),
 ("PCA vs t-SNE vs LLE?", f"PCA is linear, fast and gives loadings; t-SNE keeps local neighbours best (trustworthiness {f(DR['comparison']['t-SNE']['trustworthiness_k10'], 2)}); LLE unfolds smooth manifolds and was weakest ({f(DR['comparison']['LLE']['trustworthiness_k10'], 2)})."),
 ("How are PCA and SVD related?", "PCA is the SVD of the centred data matrix: the squared singular values give the variances. We checked this numerically."),
 ("What is trustworthiness?", "It measures whether points that are neighbours in the map were also neighbours in the original space (1 is perfect)."),
 ("Why is the genre silhouette negative?", "Genres overlap strongly in audio space; genre is partly a cultural label, so no projection separates them cleanly."),
 ("What is STL?", "Seasonal-Trend decomposition using Loess: it splits a time series into trend, seasonal and remainder. We used it on monthly release means."),
 ("What did you find about loudness?", f"Mean loudness rose from {f(fl['loudness'][0], 1)} to {f(fl['loudness'][1], 1)} dB between 1959 and 2020, with change points in 1989 and 1999."),
 ("What does the ANOVA say?", f"F = {f(T['anova']['F'], 1)}, p < 0.001: danceability differs across genres. Eta squared = {f(T['anova']['eta_squared'], 3)}, a large effect. Tukey HSD says which pairs differ."),
 ("Why also Kruskal-Wallis?", "Danceability is bounded and variances are unequal (Levene), so we confirmed with a rank-based test. It agrees."),
 ("What does the chi-square test show?", f"Mode (major/minor) depends on genre (V = {f(T['chi2_mode_genre']['cramers_v'], 2)}, moderate) but only negligibly on decade (V = {f(T['chi2_mode_decade']['cramers_v'], 2)})."),
 ("Why DuckDB-WASM?", "It runs SQL on Parquet inside the browser: fast, private, no server and free hosting on GitHub Pages."),
 ("What are the limitations?", f"Spotify features are model-made; recent years dominate; only {pct(C['join']['match_rate_pct'])} of genre tracks join to the audio table; lyrics are English-only; the 2020s have one year; the paper is not yet submitted."),
]

# ---------------------------------------------------------------- build manual (docx -> pdf) and txt
rep = Report("Manual steps for Team 05", "Team 05", "Manual steps and guide")
d = rep.doc
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("MANUAL STEPS FOR ME"); run.bold = True; run.font.size = Pt(18)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Only the jobs that need you (not Claude). Simple steps, in order.").italic = True
txt = ["MANUAL STEPS FOR ME", "Only the jobs that need you (not Claude). Simple steps, in order.", ""]
GH = L["github"]
for head, steps in SECTIONS:
    rep.h2(head)
    txt.append(head)
    for i, s in enumerate(steps, 1):
        q = d.add_paragraph(); q.paragraph_format.left_indent = Cm(0.8); q.paragraph_format.first_line_indent = Cm(-0.6)
        q.paragraph_format.space_after = Pt(3); q.paragraph_format.line_spacing = 1.2
        rr = q.add_run(f"{i}. {s}"); rr.font.size = Pt(11)
        txt.append(f"  {i}. {s}")
    txt.append("")
rep.h2("G. Last 10 points to tick before each submission")
txt.append("G. Last 10 points to tick before each submission")
for i, s in enumerate(TICK, 1):
    q = d.add_paragraph(); q.paragraph_format.left_indent = Cm(0.8); q.paragraph_format.first_line_indent = Cm(-0.6)
    q.paragraph_format.space_after = Pt(3); q.paragraph_format.line_spacing = 1.2
    q.add_run(f"[ ] {i}. {s}").font.size = Pt(11)
    txt.append(f"  [ ] {i}. {s}")
txt.append("")
rep.h1("Talking scripts", new_page=True)
txt.append("TALKING SCRIPTS")
for title, body in SCRIPTS:
    rep.h2(title)
    q = d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; q.paragraph_format.line_spacing = 1.3
    q.add_run(body).font.size = Pt(11.5)
    txt += ["", title, body]
out = OUT / f"{PF}_MANUAL_STEPS_FOR_ME.docx"
rep.save(out); to_pdf_word(out, out.with_suffix(".pdf")); out.unlink()
(OUT / f"{PF}_MANUAL_STEPS_FOR_ME.txt").write_text("\n".join(txt), encoding="utf-8")

# ---------------------------------------------------------------- viva sheet (one page)
v = Report("Viva questions - Team 05", "Team 05", "Viva questions and answers")
dv = v.doc
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(dv.sections[0], side, Cm(1.2))
p = dv.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
run = p.add_run("SonicAtlas - 20 likely viva questions (Team 05)"); run.bold = True; run.font.size = Pt(12)
for i, (qq, aa) in enumerate(VIVA, 1):
    q = dv.add_paragraph(); q.paragraph_format.space_after = Pt(1.5); q.paragraph_format.line_spacing = 1.0
    q.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = q.add_run(f"{i}. {qq} "); r1.bold = True; r1.font.size = Pt(8)
    r2 = q.add_run(aa); r2.font.size = Pt(8)
vo = OUT / f"{PF}_Viva_Questions.docx"
v.save(vo); to_pdf_word(vo, vo.with_suffix(".pdf")); vo.unlink()
import pymupdf
print("manual pages:", pymupdf.open(out.with_suffix(".pdf")).page_count, "| viva pages:", pymupdf.open(vo.with_suffix(".pdf")).page_count)
